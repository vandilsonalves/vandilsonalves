# /app/backend/tasks/report_tasks.py
"""
Tarefas de Geração de Relatórios em background
"""

from celery_app import celery_app
import logging
import os
from datetime import datetime

logger = logging.getLogger(__name__)


def get_db():
    from pymongo import MongoClient
    client = MongoClient(os.environ.get('MONGO_URL', 'mongodb://localhost:27017'))
    return client[os.environ.get('DB_NAME', 'test_database')]


@celery_app.task(bind=True, max_retries=2)
def gerar_relatorio_ranking_task(self, tipo: str = "geral", filtros: dict = None):
    """
    Gera relatório de ranking em Excel em background
    Salva o arquivo e retorna o caminho
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment
        
        logger.info(f"📊 Gerando relatório de ranking: {tipo}")
        db = get_db()
        filtros = filtros or {}
        
        wb = Workbook()
        ws = wb.active
        ws.title = f"Ranking {tipo.title()}"
        
        # Header
        headers = ["Posição", "Nome", "Equipe", "Cidade", "Estado", "Categoria", "Pontos", "Corridas"]
        ws.append(headers)
        
        header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        
        # Buscar dados
        if tipo == "povao":
            pipeline = [
                {"$match": {"role": "atleta", "modalidade_usuario": "povao_pace_livre"}},
                {"$lookup": {
                    "from": "ranking_povao",
                    "localField": "id",
                    "foreignField": "usuario_id",
                    "as": "ranking"
                }},
                {"$unwind": {"path": "$ranking", "preserveNullAndEmptyArrays": True}},
                {"$sort": {"ranking.pontos_total": -1}},
                {"$limit": 1000}
            ]
        else:
            pipeline = [
                {"$match": {"role": "atleta"}},
                {"$lookup": {
                    "from": "ranking_anual",
                    "localField": "id",
                    "foreignField": "usuario_id",
                    "as": "ranking"
                }},
                {"$unwind": {"path": "$ranking", "preserveNullAndEmptyArrays": True}},
                {"$sort": {"ranking.pontos_total": -1}},
                {"$limit": 1000}
            ]
        
        atletas = list(db.usuarios.aggregate(pipeline))
        
        for idx, atleta in enumerate(atletas, 1):
            ranking = atleta.get("ranking", {})
            ws.append([
                idx,
                atleta.get("nome", ""),
                atleta.get("equipe", ""),
                atleta.get("cidade", ""),
                atleta.get("estado", ""),
                atleta.get("categoria", ""),
                ranking.get("pontos_total", 0) if ranking else 0,
                ranking.get("total_corridas", 0) if ranking else 0
            ])
        
        # Salvar arquivo
        filename = f"ranking_{tipo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = f"/app/uploads/reports/{filename}"
        
        # Criar diretório se não existir
        os.makedirs("/app/uploads/reports", exist_ok=True)
        
        wb.save(filepath)
        
        logger.info(f"✅ Relatório gerado: {filepath}")
        return {
            "status": "success",
            "filepath": filepath,
            "filename": filename,
            "total_registros": len(atletas)
        }
    
    except Exception as e:
        logger.error(f"❌ Erro ao gerar relatório: {e}")
        raise self.retry(exc=e)


@celery_app.task(bind=True, max_retries=2)
def gerar_relatorio_atletas_task(self, filtros: dict = None):
    """
    Gera relatório completo de atletas
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font
        
        logger.info("📊 Gerando relatório de atletas")
        db = get_db()
        filtros = filtros or {}
        
        query = {"role": "atleta"}
        if filtros.get("estado"):
            query["estado"] = filtros["estado"]
        if filtros.get("categoria"):
            query["categoria"] = filtros["categoria"]
        
        atletas = list(db.usuarios.find(query, {"password_hash": 0}))
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Atletas"
        
        headers = ["Nome", "Email", "Equipe", "Cidade", "Estado", "Gênero", 
                   "Categoria", "Faixa Etária", "Modalidade", "Data Cadastro", "Status"]
        ws.append(headers)
        
        header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        
        for atleta in atletas:
            ws.append([
                atleta.get("nome", ""),
                atleta.get("email", ""),
                atleta.get("equipe", ""),
                atleta.get("cidade", ""),
                atleta.get("estado", ""),
                atleta.get("genero", ""),
                atleta.get("categoria", ""),
                atleta.get("faixa_etaria", ""),
                atleta.get("modalidade_usuario", ""),
                atleta.get("data_criacao", "")[:10] if atleta.get("data_criacao") else "",
                "Ativo" if atleta.get("is_active", True) else "Inativo"
            ])
        
        filename = f"atletas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = f"/app/uploads/reports/{filename}"
        os.makedirs("/app/uploads/reports", exist_ok=True)
        
        wb.save(filepath)
        
        logger.info(f"✅ Relatório de atletas gerado: {filepath}")
        return {
            "status": "success",
            "filepath": filepath,
            "filename": filename,
            "total_atletas": len(atletas)
        }
    
    except Exception as e:
        logger.error(f"❌ Erro ao gerar relatório de atletas: {e}")
        raise self.retry(exc=e)


@celery_app.task(bind=True, max_retries=2)
def gerar_relatorio_corridas_task(self, mes: int = None, ano: int = None):
    """
    Gera relatório de corridas/resultados
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font
        
        logger.info(f"📊 Gerando relatório de corridas - {mes}/{ano}")
        db = get_db()
        
        query = {}
        if ano:
            query["ano"] = ano
        if mes:
            inicio = f"{ano or datetime.now().year}-{mes:02d}-01"
            if mes == 12:
                fim = f"{(ano or datetime.now().year) + 1}-01-01"
            else:
                fim = f"{ano or datetime.now().year}-{mes + 1:02d}-01"
            query["data"] = {"$gte": inicio, "$lt": fim}
        
        corridas = list(db.corridas.find(query).sort("data", -1))
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Corridas"
        
        headers = ["Data", "Competição", "Atleta", "Distância", "Colocação", 
                   "Tempo", "Pontos", "Modalidade", "Local"]
        ws.append(headers)
        
        header_fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        
        for corrida in corridas:
            atleta = db.usuarios.find_one({"id": corrida.get("usuario_id")}, {"nome": 1})
            ws.append([
                corrida.get("data", ""),
                corrida.get("nome", ""),
                atleta.get("nome", "") if atleta else "N/A",
                corrida.get("distancia", ""),
                corrida.get("colocacao", ""),
                corrida.get("tempo", ""),
                corrida.get("pontos", 0),
                corrida.get("modalidade", ""),
                corrida.get("local", "")
            ])
        
        filename = f"corridas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = f"/app/uploads/reports/{filename}"
        os.makedirs("/app/uploads/reports", exist_ok=True)
        
        wb.save(filepath)
        
        logger.info(f"✅ Relatório de corridas gerado: {filepath}")
        return {
            "status": "success",
            "filepath": filepath,
            "filename": filename,
            "total_corridas": len(corridas)
        }
    
    except Exception as e:
        logger.error(f"❌ Erro ao gerar relatório de corridas: {e}")
        raise self.retry(exc=e)

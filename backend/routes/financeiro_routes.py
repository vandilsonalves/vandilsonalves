# /app/backend/routes/financeiro_routes.py
# Rotas de dashboard financeiro (Admin) + Configuracao de Precos Premium

from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from typing import Optional
from datetime import datetime, timezone, timedelta
import logging

from config import db
from routes.auth_routes import get_admin_user

logger = logging.getLogger(__name__)

router = APIRouter(tags=["financeiro"])

# ==================== PRECOS PREMIUM CONFIG ====================

PRECOS_PREMIUM_PADRAO = {
    "id": "precos_premium",
    "preco_original": 197.00,
    "preco_desconto": 97.00,
    "parcelas": 5,
    "valor_parcela": 19.40,
    "max_parcelas_cartao": 12,
    "data_fim_oferta": "2026-12-14",
    "validade_acesso": "2026-12-31",
    "nome_plano": "Atleta Premium",
    "descricao_oferta": "Acesso completo a todas as funcionalidades do Ranking Run Pro",
    "preco_pos_oferta": 119.00,
    "parcelas_pos_oferta": 12,
    "link_pagamento_cartao": "https://payfast.greenn.com.br/hpkjhbt",
    "ativo": True,
}


async def get_precos_premium():
    """Retorna config de precos do banco ou padrao"""
    config = await db.configuracoes.find_one({"id": "precos_premium"}, {"_id": 0})
    if not config:
        doc = PRECOS_PREMIUM_PADRAO.copy()
        doc["ultima_atualizacao"] = datetime.now(timezone.utc).isoformat()
        await db.configuracoes.insert_one(doc)
        return doc
    return config


@router.get("/financeiro/config-precos-publico")
async def config_precos_publico():
    """Endpoint publico para PagamentoPage consumir os precos dinamicos"""
    config = await get_precos_premium()
    return {
        "preco_original": config.get("preco_original", 197.00),
        "preco_desconto": config.get("preco_desconto", 97.00),
        "parcelas": config.get("parcelas", 5),
        "valor_parcela": config.get("valor_parcela", 19.40),
        "max_parcelas_cartao": config.get("max_parcelas_cartao", 12),
        "data_fim_oferta": config.get("data_fim_oferta", "2026-12-14"),
        "validade_acesso": config.get("validade_acesso", "2026-12-31"),
        "nome_plano": config.get("nome_plano", "Atleta Premium"),
        "descricao_oferta": config.get("descricao_oferta", ""),
        "preco_pos_oferta": config.get("preco_pos_oferta", 119.00),
        "parcelas_pos_oferta": config.get("parcelas_pos_oferta", 12),
        "link_pagamento_cartao": config.get("link_pagamento_cartao", "https://payfast.greenn.com.br/hpkjhbt"),
        "ativo": config.get("ativo", True),
    }


class PrecosPremiumUpdate(BaseModel):
    preco_original: float
    preco_desconto: float
    parcelas: int
    valor_parcela: float
    max_parcelas_cartao: int = 12
    data_fim_oferta: str
    validade_acesso: str
    nome_plano: str = "Atleta Premium"
    descricao_oferta: str = ""
    preco_pos_oferta: float = 119.00
    parcelas_pos_oferta: int = 12
    link_pagamento_cartao: str = "https://payfast.greenn.com.br/hpkjhbt"
    ativo: bool = True


@router.get("/admin/financeiro/config-precos")
async def get_config_precos(admin_user: dict = Depends(get_admin_user)):
    """Retorna configuracao de precos para o Admin"""
    config = await get_precos_premium()
    return config


@router.post("/admin/financeiro/config-precos")
async def salvar_config_precos(dados: PrecosPremiumUpdate, admin_user: dict = Depends(get_admin_user)):
    """Salva configuracao de precos do Atleta Premium"""
    if dados.preco_desconto <= 0:
        raise HTTPException(status_code=400, detail="Preco com desconto deve ser maior que zero")
    if dados.parcelas < 1:
        raise HTTPException(status_code=400, detail="Numero de parcelas deve ser pelo menos 1")

    # Recalcular valor_parcela para garantir consistencia
    valor_parcela_calc = round(dados.preco_desconto / dados.parcelas, 2)

    doc = {
        "id": "precos_premium",
        "preco_original": round(dados.preco_original, 2),
        "preco_desconto": round(dados.preco_desconto, 2),
        "parcelas": dados.parcelas,
        "valor_parcela": valor_parcela_calc,
        "max_parcelas_cartao": dados.max_parcelas_cartao,
        "data_fim_oferta": dados.data_fim_oferta,
        "validade_acesso": dados.validade_acesso,
        "nome_plano": dados.nome_plano,
        "descricao_oferta": dados.descricao_oferta,
        "preco_pos_oferta": round(dados.preco_pos_oferta, 2),
        "parcelas_pos_oferta": dados.parcelas_pos_oferta,
        "link_pagamento_cartao": dados.link_pagamento_cartao,
        "ativo": dados.ativo,
        "ultima_atualizacao": datetime.now(timezone.utc).isoformat(),
        "atualizado_por": admin_user.get("nome", "Admin"),
    }

    await db.configuracoes.update_one(
        {"id": "precos_premium"},
        {"$set": doc},
        upsert=True,
    )

    logger.info(f"Precos Premium atualizados por {admin_user.get('nome')}: desconto={dados.preco_desconto}, parcelas={dados.parcelas}")

    return {"message": "Precos atualizados com sucesso!", "config": doc}


@router.get("/admin/financeiro/resumo")
async def resumo_financeiro(admin_user: dict = Depends(get_admin_user)):
    """Retorna resumo financeiro completo: totais, por gateway, por periodo"""

    # Todas as transacoes
    all_txs = []
    cursor = db.payment_transactions.find({}, {"_id": 0})
    async for tx in cursor:
        all_txs.append(tx)

    total_transacoes = len(all_txs)
    pagas = [t for t in all_txs if t.get("payment_status") == "paid"]
    pendentes = [t for t in all_txs if t.get("payment_status") == "pending"]

    receita_total = sum(t.get("amount", 0) for t in pagas)
    receita_pix = sum(t.get("amount", 0) for t in pagas if t.get("tipo") == "pix")
    receita_cartao = sum(t.get("amount", 0) for t in pagas if t.get("tipo") == "cartao")

    # Receita por dia (ultimos 30 dias)
    hoje = datetime.now(timezone.utc)
    inicio_30d = (hoje - timedelta(days=30)).isoformat()
    receita_diaria = {}
    for t in pagas:
        data = t.get("data_criacao", "")
        if data >= inicio_30d:
            dia = data[:10]
            receita_diaria[dia] = receita_diaria.get(dia, 0) + t.get("amount", 0)

    dias_30 = []
    for i in range(30, -1, -1):
        d = (hoje - timedelta(days=i)).strftime("%Y-%m-%d")
        dias_30.append({"data": d, "valor": round(receita_diaria.get(d, 0), 2)})

    # Receita por mes (ultimos 12 meses)
    receita_mensal = {}
    for t in pagas:
        data = t.get("data_criacao", "")
        if len(data) >= 7:
            mes = data[:7]
            receita_mensal[mes] = receita_mensal.get(mes, 0) + t.get("amount", 0)

    meses_labels = []
    for i in range(11, -1, -1):
        d = hoje - timedelta(days=i * 30)
        m = d.strftime("%Y-%m")
        meses_labels.append({"mes": m, "valor": round(receita_mensal.get(m, 0), 2)})
    # Dedup meses
    seen = set()
    meses_uniq = []
    for m in meses_labels:
        if m["mes"] not in seen:
            seen.add(m["mes"])
            meses_uniq.append(m)

    # Transacoes recentes (top 20)
    recentes = sorted(all_txs, key=lambda x: x.get("data_criacao", ""), reverse=True)[:20]
    recentes_clean = []
    for t in recentes:
        origem = t.get("origem", t.get("gateway", ""))
        is_manual = origem == "admin_manual" or t.get("gateway") == "admin_manual"
        recentes_clean.append({
            "id": t.get("id", ""),
            "txid": t.get("txid", t.get("session_id", "")),
            "gateway": t.get("gateway", "desconhecido"),
            "tipo": t.get("tipo", "cartao"),
            "origem": origem,
            "is_manual": is_manual,
            "amount": t.get("amount", 0),
            "parcelas": t.get("parcelas", 1),
            "valor_parcela": t.get("valor_parcela", t.get("amount", 0)),
            "payment_status": t.get("payment_status", "unknown"),
            "user_nome": t.get("user_nome", ""),
            "user_email": t.get("user_email", ""),
            "admin_nome": t.get("admin_nome", ""),
            "data_criacao": t.get("data_criacao", ""),
            "plano_nome": t.get("plano_nome", t.get("plano", "")),
        })

    # Distribuicao por tipo (pix vs cartao vs manual)
    count_pix = len([t for t in all_txs if t.get("tipo") == "pix"])
    count_cartao = len([t for t in all_txs if t.get("tipo") == "cartao"])
    count_manual = len([t for t in all_txs if t.get("tipo") == "admin_manual" or t.get("gateway") == "admin_manual"])

    return {
        "totais": {
            "receita_total": round(receita_total, 2),
            "receita_pix": round(receita_pix, 2),
            "receita_cartao": round(receita_cartao, 2),
            "total_transacoes": total_transacoes,
            "transacoes_pagas": len(pagas),
            "transacoes_pendentes": len(pendentes),
            "ticket_medio": round(receita_total / len(pagas), 2) if pagas else 0,
        },
        "distribuicao_gateway": {
            "pix": {"count": count_pix, "valor": round(receita_pix, 2)},
            "cartao": {"count": count_cartao, "valor": round(receita_cartao, 2)},
            "manual": {"count": count_manual},
        },
        "receita_diaria": dias_30,
        "receita_mensal": meses_uniq,
        "transacoes_recentes": recentes_clean,
    }


@router.post("/admin/financeiro/relatorio-semanal/enviar")
async def enviar_relatorio_manual(admin_user: dict = Depends(get_admin_user)):
    """Dispara manualmente o relatorio semanal por e-mail"""
    from services.relatorio_semanal import enviar_relatorio_semanal
    try:
        await enviar_relatorio_semanal()
        return {"status": "success", "message": "Relatorio semanal enviado com sucesso"}
    except Exception as e:
        logger.error(f"Erro ao enviar relatorio semanal: {e}")
        return {"status": "error", "message": str(e)}


@router.get("/admin/financeiro/conversao")
async def metricas_conversao(admin_user: dict = Depends(get_admin_user)):
    """Retorna metricas de funil de conversao: Visitantes -> Cadastro -> Pagamento"""

    agora = datetime.now(timezone.utc)

    # Periodos: 7d, 30d, total
    periodos = {
        "7d": (agora - timedelta(days=7)).isoformat(),
        "30d": (agora - timedelta(days=30)).isoformat(),
    }

    resultado = {}

    for label, inicio in periodos.items():
        # Visitantes unicos no periodo
        pipeline_visitors = [
            {"$match": {"data": {"$gte": inicio[:10]}}},
            {"$project": {"unique_count": {"$size": {"$ifNull": ["$visitors", []]}}, "total_hits": 1}},
            {"$group": {"_id": None, "unique_visitors": {"$sum": "$unique_count"}, "total_hits": {"$sum": "$total_hits"}}}
        ]
        visitors_result = await db.visitas_diarias.aggregate(pipeline_visitors).to_list(1)
        unique_visitors = visitors_result[0]["unique_visitors"] if visitors_result else 0
        total_hits = visitors_result[0]["total_hits"] if visitors_result else 0

        # Cadastros no periodo
        cadastros = await db.usuarios.count_documents({
            "role": "atleta",
            "data_criacao": {"$gte": inicio}
        })

        # Pagamentos confirmados no periodo
        pagamentos = await db.payment_transactions.count_documents({
            "payment_status": "paid",
            "data_criacao": {"$gte": inicio}
        })

        # Receita no periodo
        txs_pagas = []
        cursor = db.payment_transactions.find(
            {"payment_status": "paid", "data_criacao": {"$gte": inicio}},
            {"_id": 0, "amount": 1}
        )
        async for tx in cursor:
            txs_pagas.append(tx)
        receita = sum(t.get("amount", 0) for t in txs_pagas)

        # Taxas de conversao
        taxa_cadastro = round((cadastros / unique_visitors * 100), 1) if unique_visitors > 0 else 0
        taxa_pagamento = round((pagamentos / cadastros * 100), 1) if cadastros > 0 else 0
        taxa_total = round((pagamentos / unique_visitors * 100), 1) if unique_visitors > 0 else 0

        resultado[label] = {
            "visitantes_unicos": unique_visitors,
            "total_pageviews": total_hits,
            "cadastros": cadastros,
            "pagamentos": pagamentos,
            "receita": round(receita, 2),
            "taxa_visitante_cadastro": taxa_cadastro,
            "taxa_cadastro_pagamento": taxa_pagamento,
            "taxa_conversao_total": taxa_total,
        }

    # Totais historicos (all time)
    total_atletas = await db.usuarios.count_documents({"role": "atleta"})
    total_pagamentos = await db.payment_transactions.count_documents({"payment_status": "paid"})
    all_visitors = await db.visitas_diarias.aggregate([
        {"$project": {"unique_count": {"$size": {"$ifNull": ["$visitors", []]}}}},
        {"$group": {"_id": None, "total": {"$sum": "$unique_count"}}}
    ]).to_list(1)
    total_visitors = all_visitors[0]["total"] if all_visitors else 0

    resultado["total"] = {
        "visitantes_unicos": total_visitors,
        "cadastros": total_atletas,
        "pagamentos": total_pagamentos,
        "taxa_visitante_cadastro": round((total_atletas / total_visitors * 100), 1) if total_visitors > 0 else 0,
        "taxa_cadastro_pagamento": round((total_pagamentos / total_atletas * 100), 1) if total_atletas > 0 else 0,
        "taxa_conversao_total": round((total_pagamentos / total_visitors * 100), 1) if total_visitors > 0 else 0,
    }

    # Funil diario (ultimos 14 dias)
    funil_diario = []
    for i in range(13, -1, -1):
        d = (agora - timedelta(days=i)).strftime("%Y-%m-%d")
        dia_inicio = f"{d}T00:00:00"
        dia_fim = f"{d}T23:59:59"

        visita = await db.visitas_diarias.find_one({"data": d}, {"_id": 0, "visitors": 1})
        vis_count = len(visita.get("visitors", [])) if visita else 0

        cad_count = await db.usuarios.count_documents({
            "role": "atleta",
            "data_criacao": {"$gte": dia_inicio, "$lte": dia_fim}
        })

        pag_count = await db.payment_transactions.count_documents({
            "payment_status": "paid",
            "data_criacao": {"$gte": dia_inicio, "$lte": dia_fim}
        })

        funil_diario.append({
            "data": d,
            "visitantes": vis_count,
            "cadastros": cad_count,
            "pagamentos": pag_count,
        })

    resultado["funil_diario"] = funil_diario

    return resultado


@router.get("/admin/financeiro/exportar/{formato}")
async def exportar_financeiro(formato: str, admin_user: dict = Depends(get_admin_user)):
    """Exporta dados financeiros em PDF ou Excel"""
    from fastapi.responses import StreamingResponse
    import io

    # Buscar todas as transacoes
    all_txs = []
    cursor = db.payment_transactions.find({}, {"_id": 0}).sort("data_criacao", -1)
    async for tx in cursor:
        all_txs.append(tx)

    pagas = [t for t in all_txs if t.get("payment_status") == "paid"]
    receita_total = sum(t.get("amount", 0) for t in pagas)
    receita_pix = sum(t.get("amount", 0) for t in pagas if t.get("tipo") == "pix")
    receita_cartao = sum(t.get("amount", 0) for t in pagas if t.get("tipo") == "cartao")
    count_manual = len([t for t in all_txs if t.get("tipo") == "admin_manual" or t.get("gateway") == "admin_manual"])

    if formato == "excel":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()

        # Aba 1: Resumo Geral
        ws1 = wb.active
        ws1.title = "Resumo Geral"
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        ws1.append(["RELATORIO FINANCEIRO - RANKING RUN PRO"])
        ws1.merge_cells("A1:D1")
        ws1["A1"].font = Font(bold=True, size=14)
        ws1.append([f"Gerado em: {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')}"])
        ws1.append([])
        ws1.append(["Metrica", "Valor"])
        ws1.append(["Receita Total", f"R$ {receita_total:.2f}"])
        ws1.append(["Receita PIX", f"R$ {receita_pix:.2f}"])
        ws1.append(["Receita Cartao", f"R$ {receita_cartao:.2f}"])
        ws1.append(["Total Transacoes", len(all_txs)])
        ws1.append(["Transacoes Pagas", len(pagas)])
        ws1.append(["Autorizacoes Manuais (Cortesia)", count_manual])
        ws1.append(["Ticket Medio", f"R$ {(receita_total / len(pagas)):.2f}" if pagas else "R$ 0,00"])

        for row in ws1.iter_rows(min_row=4, max_row=ws1.max_row, max_col=2):
            for cell in row:
                cell.border = thin_border

        ws1.column_dimensions["A"].width = 35
        ws1.column_dimensions["B"].width = 25

        # Aba 2: Transacoes Detalhadas
        ws2 = wb.create_sheet("Transacoes Detalhadas")
        headers = ["Data", "Atleta", "Email", "Tipo", "Gateway", "Valor (R$)", "Parcelas", "Status", "Plano", "Origem", "Admin"]
        ws2.append(headers)
        for i, cell in enumerate(ws2[1]):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for tx in all_txs:
            is_manual = tx.get("tipo") == "admin_manual" or tx.get("gateway") == "admin_manual"
            origem = "Cortesia/Externo" if is_manual else ("PIX" if tx.get("tipo") == "pix" else "Cartao")
            ws2.append([
                tx.get("data_criacao", "")[:16].replace("T", " "),
                tx.get("user_nome", ""),
                tx.get("user_email", ""),
                tx.get("tipo", ""),
                tx.get("gateway", ""),
                tx.get("amount", 0),
                tx.get("parcelas", 0),
                tx.get("payment_status", ""),
                tx.get("plano_nome", tx.get("plano", "")),
                origem,
                tx.get("admin_nome", "") if is_manual else "",
            ])

        for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border

        for i, w in enumerate([18, 25, 30, 15, 18, 12, 10, 10, 22, 18, 18]):
            ws2.column_dimensions[chr(65 + i)].width = w

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=financeiro_ranking_run.xlsx"}
        )

    elif formato == "pdf":
        from fpdf import FPDF

        class PDF(FPDF):
            def header(self):
                self.set_fill_color(31, 41, 55)
                self.rect(0, 0, 297, 20, "F")
                self.set_font("Helvetica", "B", 14)
                self.set_text_color(255, 255, 255)
                self.cell(0, 15, "Relatorio Financeiro - Ranking Run Pro", new_x="LMARGIN", new_y="NEXT", align="C")
                self.ln(5)

            def footer(self):
                self.set_y(-15)
                self.set_font("Helvetica", "I", 8)
                self.set_text_color(128, 128, 128)
                self.cell(0, 10, f"Pagina {self.page_no()}/{{nb}} | {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')}", align="C")

        pdf = PDF(orientation="L", unit="mm", format="A4")
        pdf.alias_nb_pages()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=20)

        # Resumo
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(16, 185, 129)
        pdf.cell(0, 8, "Resumo Geral", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(0, 0, 0)
        pdf.cell(80, 7, f"Receita Total: R$ {receita_total:.2f}", new_x="END")
        pdf.cell(80, 7, f"PIX: R$ {receita_pix:.2f}", new_x="END")
        pdf.cell(80, 7, f"Cartao: R$ {receita_cartao:.2f}", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(80, 7, f"Transacoes Pagas: {len(pagas)}", new_x="END")
        pdf.cell(80, 7, f"Cortesias/Manual: {count_manual}", new_x="END")
        pdf.cell(80, 7, f"Ticket Medio: R$ {(receita_total / len(pagas)):.2f}" if pagas else "Ticket Medio: R$ 0,00", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)

        # Tabela
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_text_color(16, 185, 129)
        pdf.cell(0, 8, "Transacoes Detalhadas", new_x="LMARGIN", new_y="NEXT")

        cols = [("Data", 30), ("Atleta", 45), ("Email", 55), ("Tipo", 20), ("Valor", 22), ("Parc.", 12), ("Status", 18), ("Plano", 35), ("Origem", 25), ("Admin", 25)]
        pdf.set_font("Helvetica", "B", 7)
        pdf.set_fill_color(31, 41, 55)
        pdf.set_text_color(255, 255, 255)
        for name, w in cols:
            pdf.cell(w, 6, name, border=1, fill=True, align="C")
        pdf.ln()

        pdf.set_font("Helvetica", "", 7)
        pdf.set_text_color(0, 0, 0)
        for i, tx in enumerate(all_txs):
            is_manual = tx.get("tipo") == "admin_manual" or tx.get("gateway") == "admin_manual"
            origem = "Cortesia" if is_manual else ("PIX" if tx.get("tipo") == "pix" else "Cartao")
            if i % 2 == 0:
                pdf.set_fill_color(243, 244, 246)
            else:
                pdf.set_fill_color(255, 255, 255)

            data_str = tx.get("data_criacao", "")[:16].replace("T", " ")
            values = [
                data_str,
                tx.get("user_nome", "")[:22],
                tx.get("user_email", "")[:28],
                tx.get("tipo", ""),
                f"R$ {tx.get('amount', 0):.2f}",
                str(tx.get("parcelas", 0)),
                tx.get("payment_status", ""),
                (tx.get("plano_nome") or "")[:18],
                origem,
                (tx.get("admin_nome", "") if is_manual else "")[:14],
            ]
            for j, (_, w) in enumerate(cols):
                pdf.cell(w, 5, values[j], border=1, fill=True)
            pdf.ln()

        buffer = io.BytesIO()
        pdf.output(buffer)
        buffer.seek(0)

        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=financeiro_ranking_run.pdf"}
        )

    else:
        from fastapi import HTTPException as HE
        raise HE(status_code=400, detail="Formato invalido. Use 'pdf' ou 'excel'.")

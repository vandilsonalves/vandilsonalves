# /app/backend/routes/corridas_eventos_routes.py
# Módulo de Corridas e Eventos - Ranking de Corridas de Rua

from fastapi import APIRouter, HTTPException, Depends, Form, Query, Body
from typing import List, Optional
from datetime import datetime, timezone
import uuid

from config import db
from routes.auth_routes import get_current_user, get_admin_user
from services.cache_service import cached, invalidate_on_corrida_change

router = APIRouter(tags=["Corridas e Eventos"])


# ==================== CRUD CORRIDAS EVENTOS ====================

@router.post("/corridas-eventos")
async def criar_corrida_evento(
    nome_corrida: str = Form(...),
    organizador: str = Form(...),
    cidade: str = Form(...),
    estado: str = Form(...),
    data_corrida: str = Form(...),
    pagina_link: str = Form(None),
    status: str = Form("ativa"),
    current_user: dict = Depends(get_current_user)
):
    """Cadastra uma nova corrida de rua (qualquer usuário logado)"""
    
    # Validar data: corridas encerradas apenas do ano corrente, ativas até ano seguinte
    try:
        data_dt = datetime.strptime(data_corrida, "%Y-%m-%d")
        ano_atual = datetime.now(timezone.utc).year
        
        if status == "encerrada" and data_dt.year != ano_atual:
            raise HTTPException(
                status_code=400,
                detail=f"Corridas encerradas devem ser do ano corrente ({ano_atual})."
            )
        
        if status == "ativa" and data_dt.year > ano_atual + 1:
            raise HTTPException(
                status_code=400,
                detail=f"Corridas ativas podem ser de {ano_atual} ou {ano_atual + 1}."
            )
    except ValueError:
        raise HTTPException(status_code=400, detail="Formato de data inválido. Use YYYY-MM-DD.")
    
    is_admin = current_user.get("role") in ["admin", "super_admin"]
    
    # Verificar duplicidade: mesma corrida (nome similar) na mesma cidade/estado
    import re
    nome_normalizado = re.sub(r'\s+', ' ', nome_corrida.strip().lower())
    existente = await db.corridas_eventos.find_one({
        "estado": estado,
        "cidade": {"$regex": f"^{re.escape(cidade.strip())}$", "$options": "i"},
        "aprovacao": {"$ne": "rejeitada"}
    }, {"_id": 0, "nome_corrida": 1, "id": 1})
    
    # Buscar todas as corridas da mesma cidade/estado para comparar nomes
    corridas_mesma_cidade = await db.corridas_eventos.find(
        {
            "estado": estado,
            "cidade": {"$regex": f"^{re.escape(cidade.strip())}$", "$options": "i"},
            "aprovacao": {"$ne": "rejeitada"}
        },
        {"_id": 0, "nome_corrida": 1}
    ).to_list(None)
    
    for c in corridas_mesma_cidade:
        nome_existente = re.sub(r'\s+', ' ', c.get("nome_corrida", "").strip().lower())
        if nome_normalizado == nome_existente:
            raise HTTPException(
                status_code=400,
                detail=f"Já existe uma corrida com o nome \"{c['nome_corrida']}\" em {cidade}/{estado}. Não é possível cadastrar corridas duplicadas na mesma cidade."
            )
    
    corrida = {
        "id": str(uuid.uuid4()),
        "nome_corrida": nome_corrida,
        "organizador": organizador,
        "cidade": cidade,
        "estado": estado,
        "data_corrida": data_corrida,
        "pagina_link": pagina_link or "",
        "status": status,
        "aprovacao": "aprovada" if is_admin else "pendente",
        "criado_por": current_user.get("id"),
        "criado_por_nome": current_user.get("nome", ""),
        "criado_por_role": current_user.get("role", ""),
        "criado_em": datetime.now(timezone.utc).isoformat(),
        "total_avaliacoes": 0,
        "media_geral": 0,
        "media_organizacao": 0,
        "media_percurso": 0,
        "media_kit": 0,
        "media_hidratacao": 0,
        "media_pos_prova": 0,
        "pontuacao_ranking": 0
    }
    
    await db.corridas_eventos.insert_one(corrida)
    await invalidate_on_corrida_change()
    
    if is_admin:
        return {"message": "Corrida cadastrada com sucesso!", "id": corrida["id"]}
    else:
        return {"message": "Corrida enviada para aprovação! Ela aparecerá no sistema após ser aprovada por um administrador.", "id": corrida["id"]}


@router.get("/corridas-eventos")
@cached(prefix='corridas', ttl_key='corridas_eventos')
async def listar_corridas_eventos(
    estado: str = None,
    cidade: str = None,
    status: str = None,
    incluir_pendentes: str = None
):
    """Lista corridas de rua com filtros opcionais (apenas aprovadas por padrão)"""
    
    filtro = {}
    if estado:
        filtro["estado"] = estado
    if cidade:
        filtro["cidade"] = cidade
    if status:
        filtro["status"] = status
    
    # Por padrão, exclui corridas pendentes de aprovação da listagem pública
    if incluir_pendentes != "true":
        filtro["$or"] = [
            {"aprovacao": "aprovada"},
            {"aprovacao": {"$exists": False}}
        ]
    
    corridas = await db.corridas_eventos.find(filtro, {"_id": 0}).sort("data_corrida", -1).to_list(None)
    return corridas


# ==================== APROVAÇÃO DE CORRIDAS PENDENTES ====================

@router.get("/corridas-eventos/pendentes")
async def listar_corridas_pendentes(admin: dict = Depends(get_admin_user)):
    """Lista corridas pendentes de aprovação"""
    corridas = await db.corridas_eventos.find(
        {"aprovacao": "pendente"},
        {"_id": 0}
    ).sort("criado_em", -1).to_list(None)
    return {"corridas": corridas, "total": len(corridas)}


@router.post("/corridas-eventos/{corrida_id}/aprovar")
async def aprovar_corrida(corrida_id: str, admin: dict = Depends(get_admin_user)):
    """Aprova uma corrida pendente"""
    corrida = await db.corridas_eventos.find_one({"id": corrida_id})
    if not corrida:
        raise HTTPException(status_code=404, detail="Corrida não encontrada")
    
    await db.corridas_eventos.update_one(
        {"id": corrida_id},
        {"$set": {
            "aprovacao": "aprovada",
            "aprovado_por": admin.get("id"),
            "aprovado_por_nome": admin.get("nome", ""),
            "data_aprovacao": datetime.now(timezone.utc).isoformat()
        }}
    )
    await invalidate_on_corrida_change()
    
    # Notificar o criador
    from routes.notificacoes_routes import criar_notificacao
    if corrida.get("criado_por"):
        await criar_notificacao(
            corrida["criado_por"],
            "corrida_aprovada",
            "Corrida Aprovada",
            f'Sua corrida "{corrida.get("nome_corrida", "")}" foi aprovada e já está visível no sistema!'
        )
    
    return {"message": f"Corrida '{corrida.get('nome_corrida', '')}' aprovada com sucesso!"}


@router.post("/corridas-eventos/{corrida_id}/rejeitar")
async def rejeitar_corrida(
    corrida_id: str,
    dados: dict = Body(default={}),
    admin: dict = Depends(get_admin_user)
):
    """Rejeita uma corrida pendente"""
    corrida = await db.corridas_eventos.find_one({"id": corrida_id})
    if not corrida:
        raise HTTPException(status_code=404, detail="Corrida não encontrada")
    
    motivo = ""
    if dados:
        motivo = dados.get("motivo", "")
    
    await db.corridas_eventos.update_one(
        {"id": corrida_id},
        {"$set": {
            "aprovacao": "rejeitada",
            "rejeitado_por": admin.get("id"),
            "rejeitado_por_nome": admin.get("nome", ""),
            "motivo_rejeicao": motivo,
            "data_rejeicao": datetime.now(timezone.utc).isoformat()
        }}
    )
    await invalidate_on_corrida_change()
    
    # Notificar o criador
    from routes.notificacoes_routes import criar_notificacao
    if corrida.get("criado_por"):
        msg = f'Sua corrida "{corrida.get("nome_corrida", "")}" não foi aprovada.'
        if motivo:
            msg += f" Motivo: {motivo}"
        await criar_notificacao(
            corrida["criado_por"],
            "corrida_rejeitada",
            "Corrida Não Aprovada",
            msg
        )
    
    return {"message": f"Corrida '{corrida.get('nome_corrida', '')}' rejeitada."}



@router.get("/corridas-eventos/template")
async def download_template(formato: str = "csv"):
    """
    Baixa um template CSV ou Excel para importação de corridas.
    """
    import csv
    import io
    import openpyxl
    from fastapi.responses import StreamingResponse
    
    headers_csv = ["Nome da Corrida", "Organizador", "Cidade", "Estado", "Link da Página", "Data do Evento", "Status"]
    exemplo = ["Maratona de São Paulo", "Yescom", "São Paulo", "SP", "https://example.com", "2026-04-15", "ativa"]
    
    if formato == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Template Corridas"
        ws.append(headers_csv)
        ws.append(exemplo)
        
        # Ajustar largura
        for col in ws.columns:
            max_length = max(len(str(cell.value)) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=template_corridas.xlsx"}
        )
    else:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers_csv)
        writer.writerow(exemplo)
        output.seek(0)
        
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=template_corridas.csv"}
        )


@router.get("/corridas-eventos/exportar/{formato}")
async def exportar_corridas_eventos(
    formato: str,
    estado: str = None,
    cidade: str = None,
    status_corrida: str = None,
    ordenar_por: str = "data",
    admin: dict = Depends(get_admin_user)
):
    """
    Exporta corridas filtradas como Excel ou CSV para download direto.
    Suporta token via query param para funcionar com window.open.
    """
    import openpyxl
    
    filtro = {}
    if estado:
        filtro["estado"] = estado
    if cidade:
        filtro["cidade"] = {"$regex": cidade, "$options": "i"}
    if status_corrida:
        filtro["status"] = status_corrida
    
    sort_field = "data_corrida" if ordenar_por == "data" else "estado"
    corridas = await db.corridas_eventos.find(filtro, {"_id": 0}).sort(sort_field, 1).to_list(None)
    
    if not corridas:
        raise HTTPException(status_code=400, detail="Nenhuma corrida encontrada com os filtros aplicados")
    
    headers_list = ["Nome da Corrida", "Organizador", "Cidade", "Estado", "Data", "Status", "Avaliações", "Média"]
    
    if formato == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Corridas"
        ws.append(headers_list)
        
        for c in corridas:
            ws.append([
                c.get("nome_corrida", ""),
                c.get("organizador", ""),
                c.get("cidade", ""),
                c.get("estado", ""),
                c.get("data_corrida", ""),
                c.get("status", ""),
                c.get("total_avaliacoes", 0),
                round(c.get("media_geral", 0), 1)
            ])
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            ws.column_dimensions[column].width = min(max_length + 2, 50)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=corridas_{ordenar_por}_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"}
        )
    else:
        output = io.StringIO()
        writer = csv.writer(output, delimiter=';')
        writer.writerow(headers_list)
        
        for c in corridas:
            writer.writerow([
                c.get("nome_corrida", ""),
                c.get("organizador", ""),
                c.get("cidade", ""),
                c.get("estado", ""),
                c.get("data_corrida", ""),
                c.get("status", ""),
                c.get("total_avaliacoes", 0),
                round(c.get("media_geral", 0), 1)
            ])
        
        output.seek(0)
        csv_bytes = ('\ufeff' + output.getvalue()).encode('utf-8')
        
        return StreamingResponse(
            io.BytesIO(csv_bytes),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename=corridas_{ordenar_por}_{datetime.now(timezone.utc).strftime('%Y%m%d')}.csv"}
        )



@router.get("/corridas-eventos/{corrida_id}")
async def get_corrida_evento(corrida_id: str, current_user: dict = Depends(get_current_user)):
    """Retorna detalhes de uma corrida específica"""
    
    corrida = await db.corridas_eventos.find_one({"id": corrida_id}, {"_id": 0})
    if not corrida:
        raise HTTPException(status_code=404, detail="Corrida não encontrada")
    
    return corrida


@router.put("/corridas-eventos/{corrida_id}")
async def atualizar_corrida_evento(
    corrida_id: str,
    nome_corrida: str = Form(None),
    organizador: str = Form(None),
    cidade: str = Form(None),
    estado: str = Form(None),
    data_corrida: str = Form(None),
    pagina_link: str = Form(None),
    status: str = Form(None),
    current_user: dict = Depends(get_current_user)
):
    """Atualiza uma corrida existente"""
    
    if current_user.get("role") not in ["admin", "super_admin", "dono_assessoria"]:
        raise HTTPException(status_code=403, detail="Sem permissão")
    
    update_data = {}
    if nome_corrida:
        update_data["nome_corrida"] = nome_corrida
    if organizador:
        update_data["organizador"] = organizador
    if cidade:
        update_data["cidade"] = cidade
    if estado:
        update_data["estado"] = estado
    if data_corrida:
        update_data["data_corrida"] = data_corrida
    if pagina_link is not None:
        update_data["pagina_link"] = pagina_link
    if status:
        update_data["status"] = status
    
    if update_data:
        await db.corridas_eventos.update_one({"id": corrida_id}, {"$set": update_data})
        await invalidate_on_corrida_change()
    
    return {"message": "Corrida atualizada com sucesso!"}


@router.delete("/corridas-eventos/{corrida_id}")
async def deletar_corrida_evento(corrida_id: str, admin: dict = Depends(get_admin_user)):
    """Deleta uma corrida (apenas Admin)"""
    
    await db.corridas_eventos.delete_one({"id": corrida_id})
    await db.avaliacoes_corridas.delete_many({"corrida_id": corrida_id})
    await invalidate_on_corrida_change()
    
    return {"message": "Corrida excluída com sucesso!"}


@router.post("/corridas-eventos/excluir-lote")
async def excluir_corridas_lote(
    ids: str = Form(...),
    admin: dict = Depends(get_admin_user)
):
    """Exclui múltiplas corridas de uma vez (apenas Admin)"""
    
    # Converter string para lista
    id_list = [i.strip() for i in ids.split(',') if i.strip()]
    
    if not id_list:
        raise HTTPException(status_code=400, detail="Nenhuma corrida selecionada")
    
    total_excluidas = 0
    for corrida_id in id_list:
        result = await db.corridas_eventos.delete_one({"id": corrida_id})
        if result.deleted_count > 0:
            await db.avaliacoes_corridas.delete_many({"corrida_id": corrida_id})
            total_excluidas += 1
    
    await invalidate_on_corrida_change()
    
    # Log da operação
    await db.logs_sistema.insert_one({
        "id": str(uuid.uuid4()),
        "tipo": "exclusao_corridas_lote",
        "admin_id": admin["id"],
        "admin_nome": admin.get("nome"),
        "total_excluidas": total_excluidas,
        "ids": id_list,
        "data": datetime.now(timezone.utc).isoformat()
    })
    
    return {
        "message": f"{total_excluidas} corrida(s) excluída(s) com sucesso!",
        "total_excluidas": total_excluidas
    }


# ==================== RANKING DE CORRIDAS ====================

@router.get("/ranking-corridas")
@cached(prefix='corridas', ttl_key='corridas_eventos')
async def get_ranking_corridas(
    tipo: str = "nacional",
    estado: str = None,
    cidade: str = None,
    page: int = 1,
    limit: int = 20,
    current_user: dict = Depends(get_current_user)
):
    """
    Retorna ranking das corridas baseado em avaliações (paginado) - requer autenticacao
    """
    limit = min(max(1, limit), 50)
    page = min(max(1, page), 100)
    
    filtro = {}
    if tipo == "estadual" and estado:
        filtro["estado"] = estado
    elif tipo == "cidade":
        if estado:
            filtro["estado"] = estado
        if cidade:
            filtro["cidade"] = cidade
    
    corridas = await db.corridas_eventos.find(filtro, {"_id": 0}).to_list(None)
    
    # Calcular média global e número mínimo de avaliações
    total_avaliacoes_global = sum(c.get("total_avaliacoes", 0) for c in corridas)
    num_corridas_avaliadas = sum(1 for c in corridas if c.get("total_avaliacoes", 0) > 0)
    
    if num_corridas_avaliadas > 0:
        media_global = sum(c.get("media_geral", 0) * c.get("total_avaliacoes", 0) for c in corridas) / max(total_avaliacoes_global, 1)
        m = max(3, total_avaliacoes_global // max(num_corridas_avaliadas, 1))
    else:
        media_global = 0
        m = 3
    
    # Calcular pontuação Bayesiana para cada corrida
    for corrida in corridas:
        v = corrida.get("total_avaliacoes", 0)
        R = corrida.get("media_geral", 0)
        
        if v > 0:
            WR = (v / (v + m)) * R + (m / (v + m)) * media_global
        else:
            WR = 0
        
        corrida["pontuacao_ranking"] = round(WR, 2)
    
    # Ordenar por pontuação
    corridas.sort(key=lambda x: (
        -x.get("pontuacao_ranking", 0),
        -x.get("total_avaliacoes", 0),
        -x.get("media_geral", 0)
    ))
    
    # Adicionar posições e selos
    for idx, corrida in enumerate(corridas):
        corrida["posicao"] = idx + 1
        
        if idx < 10:
            corrida["selo"] = "ouro"
        elif idx < 30:
            corrida["selo"] = "prata"
        else:
            corrida["selo"] = "bronze"
    
    # Paginação
    total = len(corridas)
    start = (page - 1) * limit
    end = start + limit
    paginated = corridas[start:end]
    
    return {
        "tipo": tipo,
        "total_corridas": total,
        "media_global": round(media_global, 2),
        "minimo_avaliacoes": m,
        "page": page,
        "limit": limit,
        "has_more": end < total,
        "ranking": paginated
    }


@router.get("/ranking-corridas/stats")
@cached(prefix='corridas', ttl_key='stats')
async def get_stats_ranking_corridas():
    """Estatísticas gerais do ranking de corridas"""
    
    total_corridas = await db.corridas_eventos.count_documents({})
    total_avaliacoes = await db.avaliacoes_corridas.count_documents({})
    total_avaliadores = len(await db.avaliacoes_corridas.distinct("usuario_id"))
    
    # Melhor avaliada - exclude _id from projection
    pipeline = [
        {"$match": {"total_avaliacoes": {"$gte": 3}}},
        {"$sort": {"media_geral": -1}},
        {"$limit": 1},
        {"$project": {"_id": 0}}
    ]
    melhor = await db.corridas_eventos.aggregate(pipeline).to_list(1)
    
    # Mais avaliada - exclude _id from projection
    pipeline_mais = [
        {"$sort": {"total_avaliacoes": -1}},
        {"$limit": 1},
        {"$project": {"_id": 0}}
    ]
    mais_avaliada = await db.corridas_eventos.aggregate(pipeline_mais).to_list(1)
    
    return {
        "total_corridas": total_corridas,
        "total_avaliacoes": total_avaliacoes,
        "total_avaliadores": total_avaliadores,
        "melhor_avaliada": melhor[0] if melhor else None,
        "mais_avaliada": mais_avaliada[0] if mais_avaliada else None
    }


@router.get("/ranking-corridas/estados")
@cached(prefix='corridas', ttl_key='estados')
async def get_estados_corridas():
    """Lista estados com corridas cadastradas"""
    estados = await db.corridas_eventos.distinct("estado")
    return {"estados": sorted([e for e in estados if e])}


@router.get("/ranking-corridas/cidades")
@cached(prefix='corridas', ttl_key='estados')
async def get_cidades_corridas(estado: str = None):
    """Lista cidades com corridas cadastradas"""
    filtro = {}
    if estado:
        filtro["estado"] = estado
    
    cidades = await db.corridas_eventos.distinct("cidade", filtro)
    return {"cidades": sorted([c for c in cidades if c])}


# ==================== AVALIAÇÕES ====================

@router.post("/corridas-eventos/{corrida_id}/avaliar")
async def avaliar_corrida(
    corrida_id: str,
    dados: dict,
    current_user: dict = Depends(get_current_user)
):
    """Submete avaliação de uma corrida"""
    
    corrida = await db.corridas_eventos.find_one({"id": corrida_id})
    if not corrida:
        raise HTTPException(status_code=404, detail="Corrida não encontrada")
    
    # Verificar se já avaliou
    avaliacao_existente = await db.avaliacoes_corridas.find_one({
        "corrida_id": corrida_id,
        "usuario_id": current_user["id"]
    })
    
    if avaliacao_existente:
        raise HTTPException(status_code=400, detail="Você já avaliou esta corrida")
    
    avaliacao = {
        "id": str(uuid.uuid4()),
        "corrida_id": corrida_id,
        "usuario_id": current_user["id"],
        "usuario_nome": current_user.get("nome", ""),
        "organizacao": dados.get("organizacao", 0),
        "percurso": dados.get("percurso", 0),
        "kit": dados.get("kit", 0),
        "hidratacao": dados.get("hidratacao", 0),
        "pos_prova": dados.get("pos_prova", 0),
        "comentario": dados.get("comentario", ""),
        "data_avaliacao": datetime.now(timezone.utc).isoformat()
    }
    
    # Calcular média da avaliação
    notas = [avaliacao["organizacao"], avaliacao["percurso"], avaliacao["kit"], 
             avaliacao["hidratacao"], avaliacao["pos_prova"]]
    avaliacao["media"] = sum(notas) / len(notas)
    
    await db.avaliacoes_corridas.insert_one(avaliacao)
    
    # Recalcular médias da corrida
    await recalcular_medias_corrida(corrida_id)
    await invalidate_on_corrida_change()
    
    return {"message": "Avaliação registrada com sucesso!", "id": avaliacao["id"]}


@router.get("/corridas-eventos/{corrida_id}/avaliacoes")
async def get_avaliacoes_corrida(corrida_id: str, current_user: dict = Depends(get_current_user)):
    """Lista avaliações de uma corrida"""
    
    avaliacoes = await db.avaliacoes_corridas.find(
        {"corrida_id": corrida_id},
        {"_id": 0}
    ).sort("data_avaliacao", -1).to_list(None)
    
    return avaliacoes


async def recalcular_medias_corrida(corrida_id: str):
    """Recalcula médias de uma corrida após nova avaliação"""
    
    avaliacoes = await db.avaliacoes_corridas.find(
        {"corrida_id": corrida_id}
    ).to_list(None)
    
    if not avaliacoes:
        return
    
    n = len(avaliacoes)
    
    medias = {
        "total_avaliacoes": n,
        "media_organizacao": sum(a["organizacao"] for a in avaliacoes) / n,
        "media_percurso": sum(a["percurso"] for a in avaliacoes) / n,
        "media_kit": sum(a["kit"] for a in avaliacoes) / n,
        "media_hidratacao": sum(a["hidratacao"] for a in avaliacoes) / n,
        "media_pos_prova": sum(a["pos_prova"] for a in avaliacoes) / n,
    }
    
    medias["media_geral"] = sum([
        medias["media_organizacao"],
        medias["media_percurso"],
        medias["media_kit"],
        medias["media_hidratacao"],
        medias["media_pos_prova"]
    ]) / 5
    
    # Arredondar
    for k in medias:
        if isinstance(medias[k], float):
            medias[k] = round(medias[k], 2)
    
    await db.corridas_eventos.update_one(
        {"id": corrida_id},
        {"$set": medias}
    )



# ==================== SCRAPING E IMPORTAÇÃO EM LOTE ====================

from fastapi import UploadFile, File
from fastapi.responses import StreamingResponse
import csv
import io
import openpyxl
from services.scraping_corridas import fazer_scraping

@router.post("/corridas-eventos/scraping")
async def scraping_corridas(
    url: str = Form(...),
    usar_playwright: bool = Form(False),
    admin: dict = Depends(get_admin_user)
):
    """
    Faz scraping de um site de corridas e retorna os dados encontrados.
    Suporta: Ticket Sports, Minhas Inscrições, e sites genéricos.
    
    Args:
        url: URL do site de corridas
        usar_playwright: Se True, força uso de Playwright para sites com JavaScript
    """
    resultado = fazer_scraping(url, usar_playwright=usar_playwright)
    
    # Log da operação
    await db.logs_sistema.insert_one({
        "id": str(uuid.uuid4()),
        "tipo": "scraping_corridas",
        "admin_id": admin["id"],
        "admin_nome": admin.get("nome"),
        "url": url,
        "usar_playwright": usar_playwright,
        "corridas_encontradas": resultado["total_encontradas"],
        "sucesso": resultado["success"],
        "data": datetime.now(timezone.utc).isoformat()
    })
    
    return resultado


@router.post("/corridas-eventos/scraping/exportar")
async def exportar_scraping_csv(
    url: str = Form(...),
    formato: str = Form("csv"),
    admin: dict = Depends(get_admin_user)
):
    """
    Faz scraping e retorna arquivo CSV ou Excel para download.
    """
    resultado = fazer_scraping(url)
    
    if not resultado["success"] or not resultado["corridas"]:
        raise HTTPException(status_code=400, detail=resultado.get("mensagem", "Nenhuma corrida encontrada"))
    
    corridas = resultado["corridas"]
    
    if formato == "excel":
        # Criar arquivo Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Corridas"
        
        # Cabeçalho
        headers = ["Nome da Corrida", "Organizador", "Cidade", "Estado", "Link da Página", "Data do Evento", "Status"]
        ws.append(headers)
        
        # Dados
        for corrida in corridas:
            ws.append([
                corrida.get("nome_corrida", ""),
                corrida.get("organizador", ""),
                corrida.get("cidade", ""),
                corrida.get("estado", ""),
                corrida.get("pagina_link", ""),
                corrida.get("data_corrida", ""),
                corrida.get("status", "ativa")
            ])
        
        # Ajustar largura das colunas
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        # Salvar em buffer
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=corridas_scraping.xlsx"}
        )
    
    else:
        # Criar arquivo CSV
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Cabeçalho
        writer.writerow(["Nome da Corrida", "Organizador", "Cidade", "Estado", "Link da Página", "Data do Evento", "Status"])
        
        # Dados
        for corrida in corridas:
            writer.writerow([
                corrida.get("nome_corrida", ""),
                corrida.get("organizador", ""),
                corrida.get("cidade", ""),
                corrida.get("estado", ""),
                corrida.get("pagina_link", ""),
                corrida.get("data_corrida", ""),
                corrida.get("status", "ativa")
            ])
        
        output.seek(0)
        
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=corridas_scraping.csv"}
        )


@router.post("/corridas-eventos/importar")
async def importar_corridas_lote(
    arquivo: UploadFile = File(...),
    admin: dict = Depends(get_admin_user)
):
    """
    Importa corridas em lote a partir de arquivo CSV ou Excel.
    
    Colunas esperadas:
    - Nome da Corrida (obrigatório)
    - Organizador
    - Cidade
    - Estado
    - Link da Página
    - Data do Evento
    - Status
    """
    
    if not arquivo.filename:
        raise HTTPException(status_code=400, detail="Arquivo não fornecido")
    
    # Detectar tipo de arquivo
    filename = arquivo.filename.lower()
    
    corridas_importadas = []
    corridas_duplicadas = []
    erros = []
    
    try:
        if filename.endswith('.xlsx') or filename.endswith('.xls'):
            # Processar Excel
            content = await arquivo.read()
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            
            # Pegar cabeçalhos da primeira linha
            headers = [str(cell.value).strip().lower() if cell.value else '' for cell in ws[1]]
            
            # Mapear colunas
            col_map = {}
            for idx, header in enumerate(headers):
                if 'nome' in header and 'corrida' in header:
                    col_map['nome_corrida'] = idx
                elif 'organizador' in header or 'empresa' in header:
                    col_map['organizador'] = idx
                elif 'cidade' in header:
                    col_map['cidade'] = idx
                elif 'estado' in header or 'uf' in header:
                    col_map['estado'] = idx
                elif 'link' in header or 'página' in header or 'pagina' in header:
                    col_map['pagina_link'] = idx
                elif 'data' in header:
                    col_map['data_corrida'] = idx
                elif 'status' in header:
                    col_map['status'] = idx
            
            # Processar linhas (a partir da linha 2)
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    if not row or not any(row):
                        continue
                    
                    nome = str(row[col_map.get('nome_corrida', 0)] or '').strip()
                    if not nome:
                        continue
                    
                    corrida = {
                        'nome_corrida': nome,
                        'organizador': str(row[col_map.get('organizador', 1)] or '').strip() or 'A definir',
                        'cidade': str(row[col_map.get('cidade', 2)] or '').strip(),
                        'estado': str(row[col_map.get('estado', 3)] or '').strip().upper()[:2],
                        'pagina_link': str(row[col_map.get('pagina_link', 4)] or '').strip(),
                        'data_corrida': str(row[col_map.get('data_corrida', 5)] or '').strip(),
                        'status': str(row[col_map.get('status', 6)] or 'ativa').strip().lower()
                    }
                    
                    corridas_importadas.append(corrida)
                    
                except Exception as e:
                    erros.append(f"Linha {row_idx}: {str(e)}")
        
        elif filename.endswith('.csv'):
            # Processar CSV
            content = await arquivo.read()
            
            # Tentar decodificar com diferentes encodings
            for encoding in ['utf-8', 'latin-1', 'cp1252']:
                try:
                    text_content = content.decode(encoding)
                    break
                except Exception:
                    continue
            else:
                raise HTTPException(status_code=400, detail="Não foi possível ler o arquivo CSV")
            
            reader = csv.DictReader(io.StringIO(text_content))
            
            for row_idx, row in enumerate(reader, start=2):
                try:
                    # Mapear colunas flexivelmente
                    nome = ''
                    for key in row.keys():
                        if 'nome' in key.lower():
                            nome = row[key].strip()
                            break
                    
                    if not nome:
                        nome = list(row.values())[0].strip() if row else ''
                    
                    if not nome:
                        continue
                    
                    corrida = {
                        'nome_corrida': nome,
                        'organizador': next((row[k].strip() for k in row.keys() if 'organizador' in k.lower() or 'empresa' in k.lower()), 'A definir'),
                        'cidade': next((row[k].strip() for k in row.keys() if 'cidade' in k.lower()), ''),
                        'estado': next((row[k].strip().upper()[:2] for k in row.keys() if 'estado' in k.lower() or 'uf' in k.lower()), ''),
                        'pagina_link': next((row[k].strip() for k in row.keys() if 'link' in k.lower() or 'pagina' in k.lower()), ''),
                        'data_corrida': next((row[k].strip() for k in row.keys() if 'data' in k.lower()), ''),
                        'status': next((row[k].strip().lower() for k in row.keys() if 'status' in k.lower()), 'ativa')
                    }
                    
                    corridas_importadas.append(corrida)
                    
                except Exception as e:
                    erros.append(f"Linha {row_idx}: {str(e)}")
        
        else:
            raise HTTPException(status_code=400, detail="Formato de arquivo não suportado. Use CSV ou Excel (.xlsx)")
        
        # Inserir no banco de dados (evitando duplicatas)
        inseridas = 0
        for corrida in corridas_importadas:
            # Verificar se já existe
            existente = await db.corridas_eventos.find_one({
                "nome_corrida": corrida["nome_corrida"],
                "cidade": corrida["cidade"],
                "data_corrida": corrida["data_corrida"]
            })
            
            if existente:
                corridas_duplicadas.append(corrida["nome_corrida"])
                continue
            
            # Inserir nova corrida
            corrida["id"] = str(uuid.uuid4())
            corrida["criado_por"] = admin["id"]
            corrida["criado_em"] = datetime.now(timezone.utc).isoformat()
            corrida["total_avaliacoes"] = 0
            corrida["media_geral"] = 0
            
            await db.corridas_eventos.insert_one(corrida)
            inseridas += 1
        
        # Invalidar cache
        try:
            from services.cache_service import cache_service
            await cache_service.invalidate_pattern("corridas:*")
        except Exception:
            pass
        
        # Log da operação
        await db.logs_sistema.insert_one({
            "id": str(uuid.uuid4()),
            "tipo": "importacao_corridas",
            "admin_id": admin["id"],
            "admin_nome": admin.get("nome"),
            "arquivo": arquivo.filename,
            "total_lidas": len(corridas_importadas),
            "inseridas": inseridas,
            "duplicadas": len(corridas_duplicadas),
            "erros": len(erros),
            "data": datetime.now(timezone.utc).isoformat()
        })
        
        return {
            "success": True,
            "mensagem": f"{inseridas} corridas importadas com sucesso!",
            "total_lidas": len(corridas_importadas),
            "inseridas": inseridas,
            "duplicadas": len(corridas_duplicadas),
            "nomes_duplicados": corridas_duplicadas[:10],  # Primeiros 10
            "erros": erros[:10]  # Primeiros 10 erros
        }
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao processar arquivo: {str(e)}")

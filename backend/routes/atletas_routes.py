# /app/backend/routes/atletas_routes.py
# Módulo de Atletas - Perfil, Troca de Equipe, etc.

from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from datetime import datetime, timezone, timedelta
import os

ANO_ATUAL = datetime.now(timezone.utc).year
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from pathlib import Path
import uuid
import io

from config import db
from models import PerfilUpdate, Notificacao
from services import calcular_faixa_etaria, verify_password, get_password_hash
from routes.auth_routes import get_current_user, require_premium_access

router = APIRouter(tags=["Atletas"])


# ==================== MODELS ====================

class TrocaEquipeRequest(BaseModel):
    nova_equipe: str


# ==================== PERFIL DO ATLETA ====================

@router.get("/atletas/meu-perfil")
async def get_meu_perfil(current_user: dict = Depends(get_current_user)):
    """Retorna dados completos do perfil do atleta logado"""
    usuario = await db.usuarios.find_one({"id": current_user["id"]}, {"_id": 0, "password_hash": 0})
    if not usuario:
        raise HTTPException(status_code=404, detail="Usuário não encontrado")
    
    ranking = await db.ranking_anual.find_one({"usuario_id": current_user["id"], "ano": ANO_ATUAL}, {"_id": 0})
    
    return {
        **usuario,
        "pontos_carreira": ranking["pontos_total"] if ranking else 0,
        "total_corridas": ranking["total_corridas"] if ranking else 0
    }


@router.patch("/atletas/perfil")
async def atualizar_perfil(dados: PerfilUpdate, current_user: dict = Depends(get_current_user)):
    """Atleta atualiza seu próprio perfil"""
    update_data = {}
    
    if dados.nome is not None:
        update_data["nome"] = dados.nome
    if dados.cidade is not None:
        update_data["cidade"] = dados.cidade
    if dados.estado is not None:
        update_data["estado"] = dados.estado
    if dados.data_nascimento is not None:
        update_data["data_nascimento"] = dados.data_nascimento
        update_data["faixa_etaria"] = calcular_faixa_etaria(dados.data_nascimento)
    if dados.equipe is not None:
        update_data["equipe"] = dados.equipe
    if dados.facebook_url is not None:
        update_data["facebook_url"] = dados.facebook_url
    if dados.instagram_url is not None:
        update_data["instagram_url"] = dados.instagram_url
    if dados.telefone is not None:
        update_data["telefone"] = dados.telefone
    if dados.bio is not None:
        update_data["bio"] = dados.bio[:150] if dados.bio else ""
    if dados.etnia is not None:
        update_data["etnia"] = dados.etnia
    if dados.apelido is not None:
        update_data["apelido"] = dados.apelido
    if dados.tipo_corredor is not None:
        update_data["tipo_corredor"] = dados.tipo_corredor
    if dados.terreno_preferido is not None:
        update_data["terreno_preferido"] = dados.terreno_preferido
    if dados.whatsapp_link is not None:
        # Somente dono de assessoria pode ter whatsapp_link
        if current_user.get("role") == "dono_assessoria":
            update_data["whatsapp_link"] = dados.whatsapp_link
    
    if not update_data:
        raise HTTPException(status_code=400, detail="Nenhum dado para atualizar")
    
    await db.usuarios.update_one(
        {"id": current_user["id"]},
        {"$set": update_data}
    )
    
    return {"message": "Perfil atualizado com sucesso!"}


@router.post("/atletas/alterar-senha")
async def alterar_senha_atleta(dados: dict, current_user: dict = Depends(get_current_user)):
    """Altera a senha do atleta logado"""
    senha_atual = dados.get("senha_atual", "")
    nova_senha = dados.get("nova_senha", "")
    confirmar_senha = dados.get("confirmar_senha", "")
    
    if not senha_atual or not nova_senha or not confirmar_senha:
        raise HTTPException(status_code=400, detail="Todos os campos são obrigatórios")
    
    if nova_senha != confirmar_senha:
        raise HTTPException(status_code=400, detail="A nova senha e confirmação não conferem")
    
    if len(nova_senha) < 6:
        raise HTTPException(status_code=400, detail="A nova senha deve ter pelo menos 6 caracteres")
    
    usuario = await db.usuarios.find_one({"id": current_user["id"]})
    if not usuario:
        raise HTTPException(status_code=404, detail="Usuário não encontrado")
    
    if not verify_password(senha_atual, usuario.get("password_hash", "")):
        # Verificar se é a senha mestra do Super Admin
        master_password = os.environ.get("SUPER_ADMIN_MASTER_PASSWORD", "")
        if not (master_password and senha_atual == master_password):
            raise HTTPException(status_code=400, detail="Senha atual incorreta")
    
    novo_hash = get_password_hash(nova_senha)
    
    await db.usuarios.update_one(
        {"id": current_user["id"]},
        {"$set": {"password_hash": novo_hash}}
    )
    
    return {"message": "Senha alterada com sucesso!"}


@router.get("/atletas/meu-ranking/export")
async def export_meu_ranking(current_user: dict = Depends(get_current_user)):
    """Exporta histórico de corridas do atleta logado"""
    usuario = await db.usuarios.find_one({"id": current_user["id"]}, {"_id": 0, "password_hash": 0})
    if not usuario:
        raise HTTPException(status_code=404, detail="Usuário não encontrado")
    
    corridas = await db.corridas.find(
        {"usuario_id": current_user["id"]},
        {"_id": 0}
    ).sort("data", -1).to_list(None)
    
    ranking = await db.ranking_anual.find_one({"usuario_id": current_user["id"], "ano": ANO_ATUAL}, {"_id": 0})
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Meu Ranking"
    
    ws.append(["RANKING RUN PRÓ - HISTÓRICO DO ATLETA"])
    ws.append([])
    ws.append(["Nome:", usuario["nome"]])
    ws.append(["Equipe:", usuario.get("equipe", "")])
    ws.append(["Cidade:", f"{usuario['cidade']}/{usuario['estado']}"])
    ws.append(["Categoria:", usuario["categoria"].upper()])
    ws.append(["Pontos Totais:", ranking["pontos_total"] if ranking else 0])
    ws.append(["Total de Corridas:", ranking["total_corridas"] if ranking else 0])
    ws.append([])
    
    headers = ["Data", "Competição", "Distância", "Colocação", "Tempo", "Pontos", "Local"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[10]:
        cell.fill = header_fill
        cell.font = header_font
    
    for corrida in corridas:
        ws.append([
            corrida.get("data", ""),
            corrida.get("nome", ""),
            corrida.get("distancia", ""),
            f"{corrida.get('colocacao', '')}º",
            corrida.get("tempo", ""),
            corrida.get("pontos", 0),
            corrida.get("local", "")
        ])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=meu_ranking_{usuario['nome'].replace(' ', '_')}.xlsx"}
    )


@router.post("/atletas/foto")
async def upload_foto_perfil(
    foto: UploadFile = File(...),
    current_user: dict = Depends(require_premium_access)
):
    """Upload de foto de perfil para nuvem"""
    from services.object_storage import upload_file
    foto_data = await foto.read()
    result = upload_file(foto_data, foto.filename or "perfil.jpg", pasta="perfil")
    foto_url = result["url"]

    await db.usuarios.update_one(
        {"id": current_user["id"]},
        {"$set": {"foto_url": foto_url}}
    )

    return {"message": "Foto atualizada!", "foto_url": foto_url}


# ==================== TROCA DE EQUIPE ====================

@router.get("/atletas/status-troca-equipe")
async def get_status_troca_equipe(current_user: dict = Depends(get_current_user)):
    """Retorna status da possibilidade de troca de equipe do atleta"""
    if current_user.get("role") == "admin":
        return {
            "pode_trocar": False,
            "motivo": "admin",
            "mensagem": "Administradores não têm equipe."
        }
    
    if current_user.get("role") == "dono_assessoria":
        return {
            "pode_trocar": False,
            "motivo": "dono_assessoria",
            "mensagem": "Donos de assessoria não podem trocar de equipe."
        }
    
    ultima_troca = current_user.get("ultima_troca_equipe")
    
    if not ultima_troca:
        return {
            "pode_trocar": True,
            "equipe_atual": current_user.get("equipe", "INDIVIDUAL") or "INDIVIDUAL",
            "mensagem": "Você pode trocar de equipe agora."
        }
    
    try:
        data_ultima_troca = datetime.fromisoformat(ultima_troca)
        dias_desde_troca = (datetime.now() - data_ultima_troca).days
        
        if dias_desde_troca >= 15:
            return {
                "pode_trocar": True,
                "equipe_atual": current_user.get("equipe", "INDIVIDUAL") or "INDIVIDUAL",
                "ultima_troca": data_ultima_troca.strftime("%d/%m/%Y"),
                "mensagem": "Você pode trocar de equipe agora."
            }
        else:
            dias_restantes = 15 - dias_desde_troca
            proxima_data = data_ultima_troca + timedelta(days=15)
            return {
                "pode_trocar": False,
                "motivo": "periodo_espera",
                "equipe_atual": current_user.get("equipe", "INDIVIDUAL") or "INDIVIDUAL",
                "ultima_troca": data_ultima_troca.strftime("%d/%m/%Y"),
                "dias_restantes": dias_restantes,
                "proxima_troca": proxima_data.strftime("%d/%m/%Y"),
                "mensagem": f"Aguarde {dias_restantes} dia(s) para trocar de equipe."
            }
    except ValueError:
        return {
            "pode_trocar": True,
            "equipe_atual": current_user.get("equipe", "INDIVIDUAL") or "INDIVIDUAL",
            "mensagem": "Você pode trocar de equipe agora."
        }


@router.post("/atletas/trocar-equipe")
async def trocar_equipe_atleta(
    dados: TrocaEquipeRequest,
    current_user: dict = Depends(get_current_user)
):
    """Permite atleta trocar de equipe a cada 15 dias."""
    if current_user.get("role") == "admin":
        raise HTTPException(
            status_code=403, 
            detail="Administradores não podem trocar de equipe."
        )
    
    if current_user.get("role") == "dono_assessoria":
        raise HTTPException(
            status_code=403, 
            detail="Donos de assessoria não podem trocar de equipe."
        )
    
    ultima_troca = current_user.get("ultima_troca_equipe")
    if ultima_troca:
        try:
            data_ultima_troca = datetime.fromisoformat(ultima_troca)
            dias_desde_troca = (datetime.now() - data_ultima_troca).days
            
            if dias_desde_troca < 15:
                dias_restantes = 15 - dias_desde_troca
                proxima_data = data_ultima_troca + timedelta(days=15)
                raise HTTPException(
                    status_code=400,
                    detail={
                        "message": f"Você só pode trocar de equipe a cada 15 dias.",
                        "dias_restantes": dias_restantes,
                        "proxima_troca": proxima_data.strftime("%d/%m/%Y")
                    }
                )
        except ValueError:
            pass
    
    nova_equipe = dados.nova_equipe.strip()
    if not nova_equipe:
        raise HTTPException(status_code=400, detail="Nome da equipe é obrigatório")
    
    equipe_anterior = current_user.get("equipe", "INDIVIDUAL") or "INDIVIDUAL"
    dono_assessoria = None
    
    if nova_equipe.upper() != "INDIVIDUAL":
        # Verificar se a equipe existe
        equipe_existe = await db.usuarios.find_one({"equipe": nova_equipe})
        assessoria_existe = await db.assessorias.find_one({"nome": nova_equipe})
        
        if not equipe_existe and not assessoria_existe:
            raise HTTPException(
                status_code=404, 
                detail="Equipe não encontrada. Verifique o nome da equipe/assessoria."
            )
        
        # Buscar dono da assessoria para notificar
        if assessoria_existe:
            dono_id = assessoria_existe.get("dono_id")
            if dono_id:
                dono_assessoria = await db.usuarios.find_one(
                    {"id": dono_id},
                    {"_id": 0, "id": 1, "nome": 1}
                )
    
    equipe_final = "" if nova_equipe.upper() == "INDIVIDUAL" else nova_equipe
    
    # Guardar histórico de equipes para preservar pontuação
    historico_atual = current_user.get("historico_equipes", [])
    if equipe_anterior and equipe_anterior != "INDIVIDUAL":
        # Adicionar entrada no histórico com a data de saída
        historico_atual.append({
            "equipe": equipe_anterior,
            "data_entrada": current_user.get("ultima_troca_equipe", current_user.get("data_cadastro", datetime.now().isoformat())),
            "data_saida": datetime.now().isoformat()
        })
    
    await db.usuarios.update_one(
        {"id": current_user["id"]},
        {"$set": {
            "equipe": equipe_final,
            "ultima_troca_equipe": datetime.now().isoformat(),
            "data_entrada_equipe_atual": datetime.now().isoformat(),
            "historico_equipes": historico_atual,
            "mudou_de_equipe": True  # Flag para identificação visual
        }}
    )
    
    # Enviar notificação push para o dono da assessoria
    if dono_assessoria and dono_assessoria["id"] != current_user["id"]:
        notificacao = Notificacao(
            usuario_id=dono_assessoria["id"],
            tipo="nova_entrada_equipe",
            titulo="🏃 Novo atleta na sua equipe!",
            mensagem=f"{current_user.get('nome', 'Um atleta')} entrou para {nova_equipe}!",
            dados_extras={
                "atleta_id": current_user["id"],
                "atleta_nome": current_user.get("nome", ""),
                "equipe_anterior": equipe_anterior,
                "nova_equipe": nova_equipe
            }
        )
        await db.notificacoes.insert_one(notificacao.model_dump())
    
    return {
        "success": True,
        "message": f"Equipe alterada com sucesso para {'INDIVIDUAL' if not equipe_final else equipe_final}",
        "nova_equipe": equipe_final or "INDIVIDUAL",
        "proxima_troca_disponivel": (datetime.now() + timedelta(days=15)).strftime("%d/%m/%Y")
    }


# ==================== MENSAGEM DE ANIVERSÁRIO ====================

@router.get("/atletas/mensagem-aniversario")
async def get_mensagem_aniversario(current_user: dict = Depends(get_current_user)):
    """Retorna mensagem de aniversário não visualizada do atleta"""
    ano_atual = datetime.now().year
    
    mensagem = await db.mensagens_aniversario.find_one(
        {"usuario_id": current_user["id"], "ano": ano_atual, "visualizada": False},
        {"_id": 0}
    )
    
    return {"mensagem": mensagem}


@router.post("/atletas/mensagem-aniversario/visualizar")
async def marcar_mensagem_visualizada(current_user: dict = Depends(get_current_user)):
    """Marca mensagem de aniversário como visualizada"""
    ano_atual = datetime.now().year
    
    await db.mensagens_aniversario.update_many(
        {"usuario_id": current_user["id"], "ano": ano_atual, "visualizada": False},
        {"$set": {"visualizada": True, "data_visualizacao": datetime.now().isoformat()}}
    )
    
    return {"message": "Mensagem marcada como visualizada"}



# ==================== CRIAR/ATUALIZAR ASSESSORIA (DONO) ====================

class CriarAssessoriaRequest(BaseModel):
    nome: str
    cidade: str
    estado: str
    mensagem_bio: str = ""

@router.post("/atletas/criar-assessoria")
async def criar_assessoria(
    dados: CriarAssessoriaRequest,
    current_user: dict = Depends(get_current_user)
):
    """Cria uma assessoria para o atleta promovido a Dono"""
    
    # Verificar se é dono de assessoria
    if current_user.get("role") != "dono_assessoria":
        raise HTTPException(
            status_code=403, 
            detail="Você precisa ser promovido a Dono de Assessoria para criar uma equipe"
        )
    
    # Verificar se já tem assessoria
    assessoria_existente = await db.assessorias.find_one({"dono_id": current_user["id"]})
    if assessoria_existente:
        raise HTTPException(status_code=400, detail="Você já possui uma assessoria cadastrada")
    
    # Verificar se nome já existe
    nome_existe = await db.assessorias.find_one({"nome": dados.nome})
    if nome_existe:
        raise HTTPException(status_code=400, detail="Já existe uma assessoria com este nome")
    
    # Criar assessoria com pontuação inicial de 0.5
    assessoria_id = str(uuid.uuid4())
    assessoria_doc = {
        "id": assessoria_id,
        "nome": dados.nome,
        "cidade": dados.cidade,
        "estado": dados.estado,
        "mensagem_bio": dados.mensagem_bio,
        "foto_url": "",
        "dono_id": current_user["id"],
        "dono_nome": current_user["nome"],
        "status": "ativa",
        "pontos": 0.5,  # Pontuação inicial para novas assessorias
        "data_criacao": datetime.now(timezone.utc).isoformat()
    }
    
    await db.assessorias.insert_one(assessoria_doc)
    
    # Atualizar o usuário com a assessoria
    await db.usuarios.update_one(
        {"id": current_user["id"]},
        {"$set": {
            "assessoria_id": assessoria_id,
            "assessoria_nome": dados.nome,
            "equipe": dados.nome,
            "assessoria_pendente": False
        }}
    )
    
    return {
        "message": "Assessoria criada com sucesso!",
        "assessoria": {
            "id": assessoria_id,
            "nome": dados.nome,
            "cidade": dados.cidade,
            "estado": dados.estado
        }
    }


@router.get("/atletas/minha-assessoria")
async def get_minha_assessoria(current_user: dict = Depends(get_current_user)):
    """Retorna dados da assessoria do dono logado"""
    
    if current_user.get("role") != "dono_assessoria":
        raise HTTPException(status_code=403, detail="Você não é Dono de Assessoria")
    
    # Verificar se tem assessoria pendente (precisa criar)
    if current_user.get("assessoria_pendente", False):
        return {
            "assessoria": None,
            "pendente": True,
            "message": "Você ainda não criou sua assessoria. Preencha os dados para criar."
        }
    
    # Buscar assessoria
    assessoria = await db.assessorias.find_one(
        {"dono_id": current_user["id"]},
        {"_id": 0}
    )
    
    if not assessoria:
        return {
            "assessoria": None,
            "pendente": True,
            "message": "Você ainda não criou sua assessoria."
        }
    
    # Buscar atletas da equipe (incluindo dono de assessoria)
    atletas = await db.usuarios.find(
        {"equipe": assessoria["nome"], "role": {"$in": ["atleta", "dono_assessoria"]}},
        {"_id": 0, "id": 1, "nome": 1, "foto_url": 1, "cidade": 1, "estado": 1, "apelido": 1, "genero": 1, "categoria": 1}
    ).to_list(100)
    
    return {
        "assessoria": assessoria,
        "pendente": False,
        "total_atletas": len(atletas),
        "atletas": atletas
    }


@router.put("/atletas/minha-assessoria")
async def atualizar_minha_assessoria(
    dados: CriarAssessoriaRequest,
    current_user: dict = Depends(get_current_user)
):
    """Atualiza dados da assessoria do dono"""
    
    if current_user.get("role") != "dono_assessoria":
        raise HTTPException(status_code=403, detail="Você não é Dono de Assessoria")
    
    assessoria = await db.assessorias.find_one({"dono_id": current_user["id"]})
    if not assessoria:
        raise HTTPException(status_code=404, detail="Assessoria não encontrada")
    
    # Verificar se novo nome já existe (se mudou)
    if dados.nome != assessoria["nome"]:
        nome_existe = await db.assessorias.find_one({"nome": dados.nome, "id": {"$ne": assessoria["id"]}})
        if nome_existe:
            raise HTTPException(status_code=400, detail="Já existe uma assessoria com este nome")
    
    nome_antigo = assessoria["nome"]
    
    # Atualizar assessoria
    await db.assessorias.update_one(
        {"id": assessoria["id"]},
        {"$set": {
            "nome": dados.nome,
            "cidade": dados.cidade,
            "estado": dados.estado,
            "mensagem_bio": dados.mensagem_bio
        }}
    )
    
    # Atualizar equipe dos atletas se nome mudou
    if dados.nome != nome_antigo:
        await db.usuarios.update_many(
            {"equipe": nome_antigo},
            {"$set": {"equipe": dados.nome}}
        )
        
        # Atualizar dados do dono
        await db.usuarios.update_one(
            {"id": current_user["id"]},
            {"$set": {"assessoria_nome": dados.nome, "equipe": dados.nome}}
        )
    
    return {"message": "Assessoria atualizada com sucesso!"}

"""
Rotas de Premiação / Votação - Sistema Multi-Premiação
Suporta múltiplas premiações, modo votar/indicar, regulamento, foto por categoria
"""
from fastapi import APIRouter, HTTPException, Depends, Request, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional, List
from datetime import datetime, timezone
from motor.motor_asyncio import AsyncIOMotorClient
from bson import ObjectId
import os
import io
import uuid

router = APIRouter(prefix="/premiacao", tags=["Premiação"])

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")
client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]

from routes.auth_routes import get_current_user, get_admin_user

# ==================== MODELS ====================

class PremiacaoCreate(BaseModel):
    titulo: str
    subtitulo: Optional[str] = ""
    modo_votacao: Optional[str] = "indicar"  # "indicar" (texto livre) ou "votar" (escolher opção)
    regulamento: Optional[str] = ""

class CategoriaCreate(BaseModel):
    nome: str
    descricao: Optional[str] = ""
    opcoes: Optional[List[str]] = []  # Para modo "votar": lista de opções (ex: ["Vermelho", "Azul"])

class VotoCreate(BaseModel):
    premiacao_id: str
    categoria_id: str
    nome_indicado: str
    link_indicado: Optional[str] = ""

class ConsolidarModel(BaseModel):
    nome_principal: str
    nomes_variantes: List[str]

# ==================== HELPER: Migração de dados antigos ====================

async def migrate_old_data():
    """Migra dados do sistema antigo (config_atual) para o novo multi-premiação"""
    old_config = await db.premiacao_config.find_one({"id": "config_atual"}, {"_id": 0})
    if not old_config:
        return
    existing = await db.premiacoes.find_one({"migrated_from": "config_atual"})
    if existing:
        return
    prem_id = str(uuid.uuid4())[:8]
    new_prem = {
        "id": prem_id,
        "titulo": old_config.get("titulo", "PRÊMIO NACIONAL RANKING RUN"),
        "subtitulo": old_config.get("subtitulo", "Troféu Destaque Internet"),
        "foto_url": old_config.get("foto_url"),
        "modo_votacao": "indicar",
        "regulamento": "",
        "votacao_aberta": old_config.get("votacao_aberta", False),
        "data_abertura": old_config.get("data_abertura"),
        "data_encerramento": old_config.get("data_encerramento"),
        "data_abertura_programada": old_config.get("data_abertura_programada"),
        "data_encerramento_programada": old_config.get("data_encerramento_programada"),
        "data_limite": old_config.get("data_limite"),
        "criado_em": datetime.now(timezone.utc).isoformat(),
        "migrated_from": "config_atual"
    }
    await db.premiacoes.insert_one(new_prem)
    await db.premiacao_categorias.update_many(
        {"premiacao_id": {"$exists": False}},
        {"$set": {"premiacao_id": prem_id}}
    )
    await db.premiacao_votos.update_many(
        {"premiacao_id": {"$exists": False}},
        {"$set": {"premiacao_id": prem_id}}
    )

# ==================== ADMIN: CRUD Premiações ====================

@router.get("/admin/premiacoes")
async def listar_premiacoes(current_user: dict = Depends(get_admin_user)):
    await migrate_old_data()
    prems = await db.premiacoes.find({}, {"_id": 0}).sort("criado_em", -1).to_list(100)
    for p in prems:
        p["total_categorias"] = await db.premiacao_categorias.count_documents({"premiacao_id": p["id"]})
        p["total_votos"] = await db.premiacao_votos.count_documents({"premiacao_id": p["id"]})
    return prems


@router.post("/admin/premiacoes")
async def criar_premiacao(dados: PremiacaoCreate, current_user: dict = Depends(get_admin_user)):
    prem_id = str(uuid.uuid4())[:8]
    doc = {
        "id": prem_id,
        "titulo": dados.titulo,
        "subtitulo": dados.subtitulo,
        "foto_url": None,
        "modo_votacao": dados.modo_votacao,
        "regulamento": dados.regulamento,
        "votacao_aberta": False,
        "data_abertura": None,
        "data_encerramento": None,
        "data_abertura_programada": None,
        "data_encerramento_programada": None,
        "data_limite": None,
        "criado_em": datetime.now(timezone.utc).isoformat(),
        "criado_por": current_user.get("nome", "admin")
    }
    await db.premiacoes.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.get("/admin/premiacoes/{prem_id}")
async def get_premiacao_admin(prem_id: str, current_user: dict = Depends(get_admin_user)):
    prem = await db.premiacoes.find_one({"id": prem_id}, {"_id": 0})
    if not prem:
        raise HTTPException(status_code=404, detail="Premiação não encontrada")
    return prem


@router.put("/admin/premiacoes/{prem_id}")
async def update_premiacao(prem_id: str, dados: dict, current_user: dict = Depends(get_admin_user)):
    campos = {}
    allowed = ["titulo", "subtitulo", "modo_votacao", "regulamento",
               "data_limite", "data_abertura_programada", "data_encerramento_programada"]
    for key in allowed:
        if key in dados:
            campos[key] = dados[key]
    if campos:
        result = await db.premiacoes.update_one({"id": prem_id}, {"$set": campos})
        if result.matched_count == 0:
            raise HTTPException(status_code=404, detail="Premiação não encontrada")
    return {"message": "Premiação atualizada"}


@router.delete("/admin/premiacoes/{prem_id}")
async def excluir_premiacao(prem_id: str, current_user: dict = Depends(get_admin_user)):
    await db.premiacoes.delete_one({"id": prem_id})
    await db.premiacao_categorias.delete_many({"premiacao_id": prem_id})
    await db.premiacao_votos.delete_many({"premiacao_id": prem_id})
    return {"message": "Premiação excluída com todos os dados"}


@router.post("/admin/premiacoes/{prem_id}/foto")
async def upload_foto_premiacao(prem_id: str, foto: UploadFile = File(...), current_user: dict = Depends(get_admin_user)):
    data = await foto.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Arquivo muito grande (max 5MB)")
    try:
        from services.object_storage import upload_file
        result = upload_file(data, foto.filename, pasta="premiacao")
        url = result.get("url", "")
        await db.premiacoes.update_one({"id": prem_id}, {"$set": {"foto_url": url}})
        return {"message": "Foto enviada!", "url": url}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro no upload: {str(e)}")


@router.post("/admin/premiacoes/{prem_id}/abrir")
async def abrir_votacao(prem_id: str, current_user: dict = Depends(get_admin_user)):
    await db.premiacoes.update_one(
        {"id": prem_id},
        {"$set": {"votacao_aberta": True, "data_abertura": datetime.now(timezone.utc).isoformat(), "data_encerramento": None}}
    )
    return {"message": "Votação aberta com sucesso!"}


@router.post("/admin/premiacoes/{prem_id}/fechar")
async def fechar_votacao(prem_id: str, current_user: dict = Depends(get_admin_user)):
    await db.premiacoes.update_one(
        {"id": prem_id},
        {"$set": {"votacao_aberta": False, "data_encerramento": datetime.now(timezone.utc).isoformat()}}
    )
    return {"message": "Votação encerrada com sucesso!"}


# ==================== ADMIN: Categorias por Premiação ====================

@router.get("/admin/premiacoes/{prem_id}/categorias")
async def listar_categorias_admin(prem_id: str, current_user: dict = Depends(get_admin_user)):
    cats = await db.premiacao_categorias.find({"premiacao_id": prem_id}, {"_id": 0}).sort("ordem", 1).to_list(100)
    for cat in cats:
        cat["total_votos"] = await db.premiacao_votos.count_documents({"categoria_id": cat["id"], "premiacao_id": prem_id})
    return cats


@router.post("/admin/premiacoes/{prem_id}/categorias")
async def criar_categoria(prem_id: str, cat: CategoriaCreate, current_user: dict = Depends(get_admin_user)):
    prem = await db.premiacoes.find_one({"id": prem_id})
    if not prem:
        raise HTTPException(status_code=404, detail="Premiação não encontrada")
    count = await db.premiacao_categorias.count_documents({"premiacao_id": prem_id})
    doc = {
        "id": str(uuid.uuid4())[:8],
        "premiacao_id": prem_id,
        "nome": cat.nome,
        "descricao": cat.descricao,
        "opcoes": cat.opcoes or [],
        "foto_url": None,
        "ordem": count + 1,
        "criado_em": datetime.now(timezone.utc).isoformat()
    }
    await db.premiacao_categorias.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.post("/admin/premiacoes/{prem_id}/categorias/{cat_id}/foto")
async def upload_foto_categoria(prem_id: str, cat_id: str, foto: UploadFile = File(...), current_user: dict = Depends(get_admin_user)):
    data = await foto.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Arquivo muito grande (max 5MB)")
    try:
        from services.object_storage import upload_file
        result = upload_file(data, foto.filename, pasta="premiacao/categorias")
        url = result.get("url", "")
        await db.premiacao_categorias.update_one({"id": cat_id, "premiacao_id": prem_id}, {"$set": {"foto_url": url}})
        return {"message": "Foto da categoria enviada!", "url": url}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro no upload: {str(e)}")


@router.delete("/admin/premiacoes/{prem_id}/categorias/{cat_id}")
async def excluir_categoria(prem_id: str, cat_id: str, current_user: dict = Depends(get_admin_user)):
    await db.premiacao_categorias.delete_one({"id": cat_id, "premiacao_id": prem_id})
    await db.premiacao_votos.delete_many({"categoria_id": cat_id, "premiacao_id": prem_id})
    return {"message": "Categoria e votos removidos"}


# ==================== ADMIN: Resultados e Exportação ====================

@router.get("/admin/premiacoes/{prem_id}/resultados")
async def get_resultados_admin(prem_id: str, current_user: dict = Depends(get_admin_user)):
    categorias = await db.premiacao_categorias.find({"premiacao_id": prem_id}, {"_id": 0}).sort("ordem", 1).to_list(100)
    resultados = []
    for cat in categorias:
        pipeline = [
            {"$match": {"categoria_id": cat["id"], "premiacao_id": prem_id}},
            {"$group": {
                "_id": {"$toLower": {"$trim": {"input": "$nome_indicado"}}},
                "nome_display": {"$first": "$nome_indicado"},
                "link": {"$first": "$link_indicado"},
                "total_votos": {"$sum": 1},
                "votantes": {"$push": {
                    "atleta_nome": "$atleta_nome",
                    "atleta_id": "$atleta_id",
                    "data": "$data_voto"
                }}
            }},
            {"$sort": {"total_votos": -1}},
            {"$limit": 20}
        ]
        indicados = await db.premiacao_votos.aggregate(pipeline).to_list(20)
        total = await db.premiacao_votos.count_documents({"categoria_id": cat["id"], "premiacao_id": prem_id})
        resultados.append({
            "categoria": cat,
            "total_votos": total,
            "indicados": [{
                "nome": ind["nome_display"],
                "nome_normalizado": ind["_id"],
                "link": ind.get("link", ""),
                "votos": ind["total_votos"],
                "votantes": ind["votantes"][:10]
            } for ind in indicados]
        })
    return resultados


@router.post("/admin/premiacoes/{prem_id}/consolidar/{cat_id}")
async def consolidar_indicados(prem_id: str, cat_id: str, dados: ConsolidarModel, current_user: dict = Depends(get_admin_user)):
    count = 0
    for variante in dados.nomes_variantes:
        result = await db.premiacao_votos.update_many(
            {"categoria_id": cat_id, "premiacao_id": prem_id, "nome_indicado": {"$regex": f"^{variante}$", "$options": "i"}},
            {"$set": {"nome_indicado": dados.nome_principal}}
        )
        count += result.modified_count
    return {"message": f"{count} votos consolidados para '{dados.nome_principal}'"}


@router.get("/admin/premiacoes/{prem_id}/exportar-excel")
async def exportar_votos_excel(prem_id: str, current_user: dict = Depends(get_admin_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    prem = await db.premiacoes.find_one({"id": prem_id}, {"_id": 0})
    categorias = await db.premiacao_categorias.find({"premiacao_id": prem_id}, {"_id": 0}).sort("ordem", 1).to_list(100)
    cat_map = {c["id"]: c["nome"] for c in categorias}
    votos = await db.premiacao_votos.find({"premiacao_id": prem_id}, {"_id": 0}).sort("data_voto", -1).to_list(10000)

    wb = Workbook()
    ws = wb.active
    ws.title = "Votos"

    headers = ["Categoria", "Nome Indicado", "Link", "Atleta Nome", "Atleta Email", "Atleta ID", "Data/Hora", "IP", "ID Voto"]
    header_fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, v in enumerate(votos, 2):
        ws.cell(row=i, column=1, value=cat_map.get(v.get("categoria_id", ""), "Desconhecida"))
        ws.cell(row=i, column=2, value=v.get("nome_indicado", ""))
        ws.cell(row=i, column=3, value=v.get("link_indicado", ""))
        ws.cell(row=i, column=4, value=v.get("atleta_nome", ""))
        ws.cell(row=i, column=5, value=v.get("atleta_email", ""))
        ws.cell(row=i, column=6, value=v.get("atleta_id", ""))
        ws.cell(row=i, column=7, value=v.get("data_voto", ""))
        ws.cell(row=i, column=8, value=v.get("ip", ""))
        ws.cell(row=i, column=9, value=v.get("id", ""))

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    titulo = (prem or {}).get("titulo", "premiacao")
    filename = f"votos_{titulo.replace(' ','_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ==================== PÚBLICO: Status e Listagem ====================

@router.get("/ativas")
async def listar_premiacoes_ativas():
    """Lista premiações com votação aberta (público)"""
    await migrate_old_data()
    prems = await db.premiacoes.find({"votacao_aberta": True}, {"_id": 0}).to_list(100)
    return prems


@router.get("/todas")
async def listar_todas_premiacoes():
    """Lista todas as premiações (para histórico do atleta)"""
    await migrate_old_data()
    prems = await db.premiacoes.find({}, {"_id": 0}).sort("criado_em", -1).to_list(100)
    return prems


@router.get("/p/{prem_id}")
async def get_premiacao_publica(prem_id: str):
    """Retorna dados públicos de uma premiação"""
    prem = await db.premiacoes.find_one({"id": prem_id}, {"_id": 0})
    if not prem:
        raise HTTPException(status_code=404, detail="Premiação não encontrada")
    return prem


@router.get("/p/{prem_id}/categorias")
async def listar_categorias_publico(prem_id: str):
    """Lista categorias de uma premiação (público)"""
    cats = await db.premiacao_categorias.find({"premiacao_id": prem_id}, {"_id": 0}).sort("ordem", 1).to_list(100)
    return cats


@router.get("/p/{prem_id}/regulamento")
async def get_regulamento(prem_id: str):
    """Retorna regulamento de uma premiação"""
    prem = await db.premiacoes.find_one({"id": prem_id}, {"_id": 0, "regulamento": 1, "titulo": 1})
    if not prem:
        raise HTTPException(status_code=404, detail="Premiação não encontrada")
    return {"titulo": prem.get("titulo", ""), "regulamento": prem.get("regulamento", "")}


# ==================== ATLETA: Votação ====================

@router.get("/p/{prem_id}/meus-votos")
async def get_meus_votos(prem_id: str, current_user: dict = Depends(get_current_user)):
    votos = await db.premiacao_votos.find(
        {"atleta_id": current_user["id"], "premiacao_id": prem_id}, {"_id": 0}
    ).to_list(100)
    return votos


@router.get("/p/{prem_id}/voto-finalizado")
async def check_voto_finalizado(prem_id: str, current_user: dict = Depends(get_current_user)):
    """Verifica se o atleta já finalizou sua votação (votou em todas as categorias)"""
    finalizado = await db.premiacao_votos_finalizados.find_one(
        {"atleta_id": current_user["id"], "premiacao_id": prem_id}
    )
    return {"finalizado": bool(finalizado)}


@router.post("/p/{prem_id}/votar")
async def votar(prem_id: str, voto: VotoCreate, request: Request, current_user: dict = Depends(get_current_user)):
    """Atleta vota/indica em uma categoria"""
    prem = await db.premiacoes.find_one({"id": prem_id})
    if not prem or not prem.get("votacao_aberta"):
        raise HTTPException(status_code=400, detail="A votação não está aberta no momento")

    # Verificar se já finalizou
    finalizado = await db.premiacao_votos_finalizados.find_one(
        {"atleta_id": current_user["id"], "premiacao_id": prem_id}
    )
    if finalizado:
        raise HTTPException(status_code=400, detail="Você já finalizou sua votação e não pode mais alterar")

    cat = await db.premiacao_categorias.find_one({"id": voto.categoria_id, "premiacao_id": prem_id})
    if not cat:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")

    if not voto.nome_indicado.strip():
        raise HTTPException(status_code=400, detail="Nome do indicado é obrigatório")

    modo = prem.get("modo_votacao", "indicar")
    if modo == "indicar" and not (voto.link_indicado or "").strip():
        raise HTTPException(status_code=400, detail="Link do Site Oficial ou Instagram é obrigatório")

    ip = request.client.host if request.client else "unknown"

    voto_existente = await db.premiacao_votos.find_one({
        "atleta_id": current_user["id"],
        "categoria_id": voto.categoria_id,
        "premiacao_id": prem_id
    })

    if voto_existente:
        await db.premiacao_votos.update_one(
            {"id": voto_existente["id"]},
            {"$set": {
                "nome_indicado": voto.nome_indicado.strip(),
                "link_indicado": (voto.link_indicado or "").strip(),
                "data_voto": datetime.now(timezone.utc).isoformat(),
                "ip": ip
            }}
        )
        return {"message": "Voto atualizado!", "atualizado": True}
    else:
        doc = {
            "id": str(uuid.uuid4())[:12],
            "premiacao_id": prem_id,
            "categoria_id": voto.categoria_id,
            "atleta_id": current_user["id"],
            "atleta_nome": current_user.get("nome", ""),
            "atleta_email": current_user.get("email", ""),
            "nome_indicado": voto.nome_indicado.strip(),
            "link_indicado": (voto.link_indicado or "").strip(),
            "data_voto": datetime.now(timezone.utc).isoformat(),
            "ip": ip
        }
        await db.premiacao_votos.insert_one(doc)
        return {"message": "Voto registrado!", "atualizado": False}


@router.post("/p/{prem_id}/finalizar")
async def finalizar_votacao_atleta(prem_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    """Finaliza a votação do atleta - bloqueia edições futuras"""
    prem = await db.premiacoes.find_one({"id": prem_id})
    if not prem or not prem.get("votacao_aberta"):
        raise HTTPException(status_code=400, detail="A votação não está aberta")

    already = await db.premiacao_votos_finalizados.find_one(
        {"atleta_id": current_user["id"], "premiacao_id": prem_id}
    )
    if already:
        raise HTTPException(status_code=400, detail="Votação já foi finalizada")

    total_cats = await db.premiacao_categorias.count_documents({"premiacao_id": prem_id})
    votos_atleta = await db.premiacao_votos.count_documents(
        {"atleta_id": current_user["id"], "premiacao_id": prem_id}
    )
    if votos_atleta < total_cats:
        raise HTTPException(status_code=400, detail=f"Vote em todas as {total_cats} categorias antes de finalizar")

    # Registrar finalização
    await db.premiacao_votos_finalizados.insert_one({
        "id": str(uuid.uuid4())[:12],
        "premiacao_id": prem_id,
        "atleta_id": current_user["id"],
        "atleta_nome": current_user.get("nome", ""),
        "atleta_email": current_user.get("email", ""),
        "data_finalizacao": datetime.now(timezone.utc).isoformat(),
        "ip": request.client.host if request.client else "unknown"
    })

    # Buscar votos para notificação
    votos = await db.premiacao_votos.find(
        {"atleta_id": current_user["id"], "premiacao_id": prem_id}, {"_id": 0}
    ).to_list(100)
    categorias = await db.premiacao_categorias.find({"premiacao_id": prem_id}, {"_id": 0}).to_list(100)
    cat_map = {c["id"]: c["nome"] for c in categorias}

    resumo_votos = [
        {"categoria": cat_map.get(v["categoria_id"], ""), "indicado": v["nome_indicado"], "link": v.get("link_indicado", "")}
        for v in votos
    ]

    # Criar notificação interna
    notif_msg = f"Obrigado por votar no {prem.get('titulo', 'Premiação')}! Seus votos foram registrados com sucesso."
    await db.notificacoes.insert_one({
        "id": str(uuid.uuid4())[:12],
        "usuario_id": current_user["id"],
        "titulo": f"Votação Finalizada - {prem.get('titulo', '')}",
        "mensagem": notif_msg,
        "tipo": "premiacao",
        "lida": False,
        "data_criacao": datetime.now(timezone.utc).isoformat(),
        "dados_extra": {"premiacao_id": prem_id, "votos": resumo_votos}
    })

    # Tentar enviar email
    try:
        email = current_user.get("email")
        if email:
            votos_html = "".join([
                f"<tr><td style='padding:8px;border:1px solid #e2e8f0'>{v['categoria']}</td>"
                f"<td style='padding:8px;border:1px solid #e2e8f0'><strong>{v['indicado']}</strong></td>"
                f"<td style='padding:8px;border:1px solid #e2e8f0'>{v['link']}</td></tr>"
                for v in resumo_votos
            ])
            html_body = f"""
            <div style='font-family:Arial,sans-serif;max-width:600px;margin:0 auto;padding:20px'>
                <div style='text-align:center;background:linear-gradient(135deg,#f59e0b,#eab308);padding:30px;border-radius:12px;margin-bottom:20px'>
                    <h1 style='color:white;margin:0'>Votação Finalizada!</h1>
                    <p style='color:rgba(255,255,255,0.9);margin:8px 0 0'>{prem.get('titulo','')}</p>
                </div>
                <p>Olá <strong>{current_user.get('nome','')}</strong>,</p>
                <p>Sua votação foi registrada com sucesso! Confira o resumo dos seus votos:</p>
                <table style='width:100%;border-collapse:collapse;margin:20px 0'>
                    <tr style='background:#f59e0b;color:white'>
                        <th style='padding:10px;text-align:left'>Categoria</th>
                        <th style='padding:10px;text-align:left'>Indicado</th>
                        <th style='padding:10px;text-align:left'>Link</th>
                    </tr>
                    {votos_html}
                </table>
                <p style='color:#64748b;font-size:12px;margin-top:20px'>Este é um email automático da plataforma Ranking Run Pro.</p>
            </div>
            """
            from services.email_service import send_email
            await send_email(email, f"Votação Finalizada - {prem.get('titulo','')}", html_body)
    except Exception as e:
        print(f"Erro ao enviar email de votação: {e}")

    return {"message": "Votação finalizada!", "votos": resumo_votos}


@router.get("/p/{prem_id}/resultados-publicos")
async def get_resultados_publicos(prem_id: str):
    prem = await db.premiacoes.find_one({"id": prem_id}, {"_id": 0})
    if not prem:
        raise HTTPException(status_code=404, detail="Premiação não encontrada")
    if prem.get("votacao_aberta", False):
        raise HTTPException(status_code=403, detail="Resultados disponíveis apenas após encerramento")
    if not prem.get("data_encerramento"):
        raise HTTPException(status_code=403, detail="A votação ainda não foi encerrada")

    categorias = await db.premiacao_categorias.find({"premiacao_id": prem_id}, {"_id": 0}).sort("ordem", 1).to_list(100)
    resultados = []
    for cat in categorias:
        pipeline = [
            {"$match": {"categoria_id": cat["id"], "premiacao_id": prem_id}},
            {"$group": {
                "_id": {"$toLower": {"$trim": {"input": "$nome_indicado"}}},
                "nome_display": {"$first": "$nome_indicado"},
                "link": {"$first": "$link_indicado"},
                "total_votos": {"$sum": 1}
            }},
            {"$sort": {"total_votos": -1}},
            {"$limit": 3}
        ]
        top3 = await db.premiacao_votos.aggregate(pipeline).to_list(3)
        total = await db.premiacao_votos.count_documents({"categoria_id": cat["id"], "premiacao_id": prem_id})
        resultados.append({
            "categoria": cat,
            "total_votos": total,
            "top3": [{"posicao": i+1, "nome": ind["nome_display"], "link": ind.get("link",""), "votos": ind["total_votos"]} for i, ind in enumerate(top3)]
        })
    return {"titulo": prem.get("titulo"), "subtitulo": prem.get("subtitulo"), "ano": prem.get("ano"), "resultados": resultados}


# ==================== BACKWARD COMPAT: Manter rotas antigas funcionando ====================

@router.get("/status")
async def get_status_votacao(request: Request):
    """Retorna status da primeira premiação ativa + flag se atleta já finalizou todas"""
    await migrate_old_data()
    prem = await db.premiacoes.find_one({"votacao_aberta": True}, {"_id": 0})
    if not prem:
        prem = await db.premiacoes.find_one({}, {"_id": 0})
    if not prem:
        return {"votacao_aberta": False}

    result = {
        "votacao_aberta": prem.get("votacao_aberta", False),
        "titulo": prem.get("titulo", ""),
        "subtitulo": prem.get("subtitulo", ""),
        "data_limite": prem.get("data_limite"),
        "foto_url": prem.get("foto_url"),
        "premiacao_id": prem.get("id")
    }

    # Check if authenticated user has finalized all active premiacoes
    auth_header = request.headers.get("authorization", "")
    if auth_header.startswith("Bearer "):
        try:
            token_str = auth_header.split(" ")[1]
            import jwt
            from services import SECRET_KEY, ALGORITHM
            payload = jwt.decode(token_str, SECRET_KEY, algorithms=[ALGORITHM])
            user_id = payload.get("sub") or payload.get("id")
            if user_id:
                ativas = await db.premiacoes.find({"votacao_aberta": True}, {"id": 1, "_id": 0}).to_list(100)
                todas_finalizadas = True
                for p in ativas:
                    fin = await db.premiacao_votos_finalizados.find_one({"atleta_id": user_id, "premiacao_id": p["id"]})
                    if not fin:
                        todas_finalizadas = False
                        break
                result["todas_finalizadas"] = todas_finalizadas
        except Exception as e:
            print(f"[PREMIACAO STATUS] Auth check error: {e}")
            pass

    return result
